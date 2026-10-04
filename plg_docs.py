"""Document studio: PDF, Word, PowerPoint, Excel, website, invoice, data
analysis and charts."""
import datetime as dt
import html
import math
import re
import statistics
from collections import Counter
from pathlib import Path

import brand
import industries
import textkit as tk
from llm import LLMError
from registry import PluginError, plugin


def _meta(ctx) -> dict:
    return {"business": ctx.business, "signature": ctx.profile.get("signature", True) is not False,
            "date": dt.date.today().strftime("%d %B %Y")}


def _need_doc(ctx):
    doc = ctx.state.get("doc")
    if not doc:
        text = ctx.input_text()
        if text.strip():
            title = ctx.inputs[0]["name"].rsplit(".", 1)[0] if ctx.inputs else "Document"
            doc = tk.markdown_to_doc(text, tk.title_case(title.replace("_", " ")))
            ctx.state["doc"] = doc
    if not doc:
        raise PluginError("There is no content to render yet. Ask for a topic (e.g. “make a PDF about solar energy”) "
                          "or upload a document.")
    return doc


def _fname(doc, ext):
    return tk.slugify(doc.get("title", "document")) + ext


# ====================================================================== PDF
def _rl_inline(text: str) -> str:
    """Markdown-ish inline text -> well-formed reportlab markup (runs never overlap)."""
    text = brand.pdf_safe(text or "")
    runs = tk.md_runs(text) if text.count("*") <= 200 else [(text, False, False)]
    out = []
    for run, bold, italic in runs:
        t = html.escape(run, quote=False)
        t = re.sub(r"\[([^\]]+)\]\((https?://[^)\s\"<>]+)\)",
                   lambda m: f'<link href="{m.group(2)}" color="#8A6A32">{m.group(1)}</link>', t)
        if bold:
            t = f"<b>{t}</b>"
        if italic:
            t = f"<i>{t}</i>"
        out.append(t)
    return "".join(out).replace("\n", "<br/>")


def _para(text: str, style, wrap: str = ""):
    """Paragraph that can never crash the render: falls back to escaped plain text."""
    from reportlab.platypus import Paragraph
    try:
        markup = _rl_inline(text)
        return Paragraph(f"<{wrap}>{markup}</{wrap}>" if wrap else markup, style)
    except Exception:
        return Paragraph(html.escape(tk.strip_md(brand.pdf_safe(text or "")), quote=False), style)


def render_pdf(doc: dict, path, meta: dict, plain: bool = False):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (Image as RLImage, ListFlowable, ListItem, PageBreak, Paragraph,
                                    SimpleDocTemplate, Spacer, Table, TableStyle)
    from reportlab.platypus.flowables import HRFlowable

    f = brand.reportlab_fonts()
    gold, ink, gold_dark = colors.HexColor("#C9A55C"), colors.HexColor("#1A1814"), colors.HexColor("#8A6A32")
    st = {
        "title": ParagraphStyle("t", fontName=f["serif_bold"], fontSize=28, leading=34, textColor=ink, spaceAfter=10),
        "subtitle": ParagraphStyle("s", fontName=f["regular"], fontSize=13, leading=18, textColor=colors.HexColor("#5A544A")),
        "meta": ParagraphStyle("m", fontName=f["regular"], fontSize=10, leading=14, textColor=colors.HexColor("#7A7263")),
        "h1": ParagraphStyle("h1", fontName=f["serif_bold"], fontSize=17, leading=22, textColor=gold_dark,
                             spaceBefore=14, spaceAfter=6),
        "h3": ParagraphStyle("h3", fontName=f["bold"], fontSize=12, leading=16, textColor=ink, spaceBefore=8, spaceAfter=3),
        "body": ParagraphStyle("b", fontName=f["regular"], fontSize=10.5, leading=15.5, textColor=ink, spaceAfter=7),
        "quote": ParagraphStyle("q", fontName=f["regular"], fontSize=10.5, leading=15, textColor=colors.HexColor("#4A4438"),
                                leftIndent=12, borderPadding=6),
        "cell": ParagraphStyle("c", fontName=f["regular"], fontSize=9, leading=12, textColor=ink),
        "cellh": ParagraphStyle("ch", fontName=f["bold"], fontSize=9, leading=12, textColor=colors.white),
        "toc": ParagraphStyle("toc", fontName=f["regular"], fontSize=11, leading=18, textColor=ink),
    }
    W = A4[0] - 40 * mm

    def on_page(canvas, d):
        canvas.saveState()
        if plain:
            canvas.setFont(f["regular"], 8)
            canvas.setFillColor(colors.HexColor("#7A7263"))
            canvas.drawCentredString(A4[0] / 2, 10 * mm, f"{d.page}")
        elif d.page > 1:
            canvas.setStrokeColor(gold)
            canvas.setLineWidth(0.6)
            canvas.line(20 * mm, A4[1] - 14 * mm, A4[0] - 20 * mm, A4[1] - 14 * mm)
            canvas.setFont(f["regular"], 8)
            canvas.setFillColor(colors.HexColor("#7A7263"))
            canvas.drawString(20 * mm, A4[1] - 12 * mm, brand.pdf_safe(doc.get("title", ""))[:90])
            canvas.drawString(20 * mm, 10 * mm, brand.pdf_safe(meta.get("business", "")))
            canvas.drawRightString(A4[0] - 20 * mm, 10 * mm, f"Page {d.page}")
            if meta.get("signature") and brand.LOGO_MARK.exists():
                canvas.drawImage(str(brand.LOGO_MARK), A4[0] / 2 - 4 * mm, 7 * mm, width=8 * mm, height=5.6 * mm,
                                 mask="auto", preserveAspectRatio=True)
        else:
            canvas.setFillColor(gold)
            canvas.rect(0, A4[1] - 8 * mm, A4[0], 8 * mm, stroke=0, fill=1)
            canvas.rect(0, 0, A4[0], 3 * mm, stroke=0, fill=1)
        canvas.restoreState()

    story = [] if plain else [Spacer(1, 30 * mm)]
    if not plain and meta.get("signature") and brand.LOGO.exists():
        story += [RLImage(str(brand.LOGO), width=38 * mm, height=40 * mm, kind="proportional"), Spacer(1, 14 * mm)]
    story.append(_para(doc.get("title", "Document"), st["title"]))
    if doc.get("subtitle"):
        story.append(_para(doc["subtitle"], st["subtitle"]))
    if not plain:
        story += [Spacer(1, 6 * mm), HRFlowable(width="30%", thickness=2, color=gold, hAlign="LEFT"),
                  Spacer(1, 4 * mm),
                  _para(f"Prepared for {meta.get('business', '')} · {meta.get('date', '')}", st["meta"])]
        headings = [s["heading"] for s in doc.get("sections", []) if s.get("heading")]
        if len(headings) >= 4:
            story += [Spacer(1, 16 * mm), Paragraph("<b>Contents</b>", st["h3"])]
            story += [_para(f"{i}.  {h}", st["toc"]) for i, h in enumerate(headings, 1)]
        story.append(PageBreak())

    for sec in doc.get("sections", []):
        if sec.get("heading"):
            story.append(_para(sec["heading"], st["h1"]))
        for b in sec.get("blocks", []):
            t = b.get("type")
            if t == "p":
                story.append(_para(b["text"], st["body"]))
            elif t == "h3":
                story.append(_para(b["text"], st["h3"]))
            elif t == "quote":
                tbl = Table([[_para(b["text"], st["quote"])]], colWidths=[W])
                tbl.setStyle(TableStyle([("LINEBEFORE", (0, 0), (0, 0), 2.5, gold),
                                         ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FAF6EC")),
                                         ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
                story += [tbl, Spacer(1, 6)]
            elif t in ("bullets", "numbered"):
                items = [ListItem(_para(i, st["body"]), leftIndent=14) for i in b["items"]]
                story.append(ListFlowable(items, bulletType="bullet" if t == "bullets" else "1",
                                          start="•" if t == "bullets" else None, bulletColor=gold_dark,
                                          bulletFontName=f["bold"], leftIndent=14))
            elif t == "table" and b.get("rows"):
                rows = b["rows"]
                ncol = max(len(r) for r in rows)
                rows = [r + [""] * (ncol - len(r)) for r in rows]
                data = [[_para(c, st["cellh"] if ri == 0 else st["cell"]) for c in r]
                        for ri, r in enumerate(rows)]
                if ncol == 2:
                    widths = [W * 0.3, W * 0.7]
                else:
                    widths = [W / ncol] * ncol
                tbl = Table(data, colWidths=widths, repeatRows=1)
                tbl.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1A2540")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FAF6EC")]),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9CDB2")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
                story += [tbl, Spacer(1, 8)]
            elif t == "image" and b.get("path") and Path(b["path"]).exists():
                story += [RLImage(b["path"], width=W, height=W * 0.6, kind="proportional"), Spacer(1, 4)]
                if b.get("caption"):
                    story.append(_para(b["caption"], st["meta"], "i"))
    if doc.get("sources"):
        story.append(Paragraph("Sources &amp; Attribution", st["h1"]))
        for s in doc["sources"]:
            line = s.get("title", "")
            if s.get("url"):
                line += f" — [{s['url']}]({s['url']})"
            if s.get("license"):
                line += f" ({s['license']})"
            story.append(_para(line, st["meta"]))
    pdf = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=20 * mm,
                            bottomMargin=18 * mm, title=brand.pdf_safe(doc.get("title", "")),
                            author=brand.pdf_safe(meta.get("business", "")), creator="InEra All-in-One Plugins")
    pdf.build(story, onFirstPage=on_page, onLaterPages=on_page)


@plugin("pdf_creator", "PDF Creator", "documents", "render",
        "Turns any generated or uploaded content into a polished, branded PDF with cover, contents and sources.",
        icon="file-pdf", keywords=["pdf", "report", "ebook", "handout", "brochure"],
        example="Create a PDF guide on electric vehicle maintenance")
def pdf_creator(ctx):
    doc = _need_doc(ctx)
    path = ctx.out_path(_fname(doc, ".pdf"))
    render_pdf(doc, path, _meta(ctx))
    ctx.add_output(path, "PDF document")
    ctx.log("PDF created.")


# ====================================================================== Word
def _add_runs(par, text):
    for run_text, bold, italic in tk.md_runs(text):
        r = par.add_run(run_text)
        r.bold, r.italic = bold or None, italic or None


def _page_number_field(paragraph):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def render_docx(doc: dict, path, meta: dict):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    d = Document()
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.9)
    normal = d.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for name, size, color in (("Title", 30, (0x1A, 0x18, 0x14)), ("Heading 1", 18, (0x8A, 0x6A, 0x32)),
                              ("Heading 2", 14, (0x8A, 0x6A, 0x32)), ("Heading 3", 12, (0x1A, 0x18, 0x14))):
        s = d.styles[name]
        s.font.name = "Georgia"
        s.font.size = Pt(size)
        s.font.color.rgb = RGBColor(*color)
    if meta.get("signature") and brand.LOGO.exists():
        d.add_picture(str(brand.LOGO), width=Inches(1.4))
        d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.LEFT
    d.add_paragraph(doc.get("title", "Document"), style="Title")
    if doc.get("subtitle"):
        p = d.add_paragraph()
        r = p.add_run(doc["subtitle"])
        r.italic = True
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(0x5A, 0x54, 0x4A)
    p = d.add_paragraph()
    r = p.add_run(f"Prepared for {meta.get('business', '')} · {meta.get('date', '')}")
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x7A, 0x72, 0x63)
    d.add_page_break()
    for s in doc.get("sections", []):
        if s.get("heading"):
            d.add_heading(s["heading"], level=1)
        for b in s.get("blocks", []):
            t = b.get("type")
            if t == "p":
                _add_runs(d.add_paragraph(), b["text"])
            elif t == "h3":
                d.add_heading(tk.strip_md(b["text"]), level=3)
            elif t == "quote":
                _add_runs(d.add_paragraph(style="Intense Quote"), b["text"])
            elif t in ("bullets", "numbered"):
                for item in b["items"]:
                    _add_runs(d.add_paragraph(style="List Bullet" if t == "bullets" else "List Number"), item)
            elif t == "table" and b.get("rows"):
                rows = b["rows"]
                ncol = max(len(r) for r in rows)
                table = d.add_table(rows=len(rows), cols=ncol)
                table.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci in range(ncol):
                        cell = table.cell(ri, ci)
                        cell.text = ""
                        _add_runs(cell.paragraphs[0], row[ci] if ci < len(row) else "")
                        if ri == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
                                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                            shd = OxmlElement("w:shd")
                            shd.set(qn("w:val"), "clear")
                            shd.set(qn("w:color"), "auto")
                            shd.set(qn("w:fill"), "1A2540")
                            cell._tc.get_or_add_tcPr().append(shd)
                d.add_paragraph()
            elif t == "image" and b.get("path") and Path(b["path"]).exists():
                d.add_picture(b["path"], width=Inches(6.0))
                if b.get("caption"):
                    cap = d.add_paragraph()
                    cr = cap.add_run(b["caption"])
                    cr.italic = True
    if doc.get("sources"):
        d.add_heading("Sources & Attribution", level=1)
        for s in doc["sources"]:
            d.add_paragraph(" — ".join(x for x in (s.get("title"), s.get("url"), s.get("license")) if x),
                            style="List Bullet")
    footer = d.sections[0].footer.paragraphs[0]
    footer.text = f"{meta.get('business', '')}   ·   Page "
    _page_number_field(footer)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.core_properties.title = doc.get("title", "")
    d.core_properties.author = meta.get("business", "")
    d.save(str(path))


@plugin("word_creator", "Word Document", "documents", "render",
        "Creates an editable Word (.docx) document with headings, lists, tables, page numbers and sources.",
        icon="file-word", keywords=["word", "docx", "doc file", "editable document", "ms word"],
        example="Write a business proposal for a car detailing studio as a Word file")
def word_creator(ctx):
    doc = _need_doc(ctx)
    path = ctx.out_path(_fname(doc, ".docx"))
    render_docx(doc, path, _meta(ctx))
    ctx.add_output(path, "Word document")
    ctx.log("Word document created.")


# ====================================================================== Slides
def slides_from_doc(doc: dict, max_slides: int = 0) -> list:
    slides = [{"kind": "title", "title": doc.get("title", "Presentation"), "subtitle": doc.get("subtitle", "")}]
    sections = [s for s in doc.get("sections", []) if s.get("heading") or s.get("blocks")]
    if len(sections) >= 3:
        slides.append({"kind": "content", "title": "Agenda",
                       "bullets": [tk.strip_md(s["heading"]) for s in sections if s.get("heading")][:7],
                       "notes": "Here is what we will cover."})
    for s in sections:
        bullets, notes, table = [], [], None
        for b in s.get("blocks", []):
            if b["type"] in ("bullets", "numbered"):
                bullets += [tk.truncate_words(tk.strip_md(i), 22) for i in b["items"]]
                notes += [tk.strip_md(i) for i in b["items"]]
            elif b["type"] in ("p", "quote"):
                text = tk.strip_md(b["text"])
                notes.append(text)
                sents = tk.split_sentences(text)
                pick = tk.summarize(text, 3) if len(sents) > 3 else sents
                bullets += [tk.truncate_words(x, 22) for x in pick]
            elif b["type"] == "table" and b.get("rows") and table is None:
                table = b["rows"][:9]
        title = tk.strip_md(s.get("heading") or doc.get("title", ""))
        if table:
            slides.append({"kind": "table", "title": title, "table": table, "notes": " ".join(notes)})
        chunks = [bullets[i:i + 5] for i in range(0, len(bullets), 5)] or ([] if table else [[title]])
        for ci, chunk in enumerate(chunks[:3]):
            slides.append({"kind": "content", "title": title + (" (cont.)" if ci else ""), "bullets": chunk,
                           "notes": " ".join(notes) if ci == 0 else ""})
    slides.append({"kind": "end", "title": "Thank you", "subtitle": doc.get("title", "")})
    if max_slides and len(slides) > max_slides:
        keep = slides[:max_slides - 1] + [slides[-1]]
        slides = keep
    return slides


def render_pptx(doc: dict, path, meta: dict, max_slides: int = 0):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    INK, GOLD, CREAM, MUTED = (RGBColor(0x14, 0x1D, 0x31), RGBColor(0xC9, 0xA5, 0x5C), RGBColor(0xF5, 0xEF, 0xE2),
                               RGBColor(0xA3, 0xAB, 0xBC))
    plan = slides_from_doc(doc, max_slides)
    total = len(plan)

    def base(slide):
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = INK
        bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.08))
        bar.fill.solid()
        bar.fill.fore_color.rgb = GOLD
        bar.line.fill.background()

    def textbox(slide, x, y, w, h, text, size, color, bold=False, font="Georgia", align=None):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size, p.font.bold, p.font.name = Pt(size), bold, font
        p.font.color.rgb = color
        if align:
            p.alignment = align
        return tb

    for i, s in enumerate(plan, 1):
        slide = prs.slides.add_slide(blank)
        base(slide)
        if s["kind"] == "title":
            if meta.get("signature") and brand.LOGO.exists():
                slide.shapes.add_picture(str(brand.LOGO), Inches(0.9), Inches(0.8), height=Inches(1.9))
            textbox(slide, Inches(0.9), Inches(3.2), Inches(11.5), Inches(1.8), s["title"], 40, CREAM, True)
            if s.get("subtitle"):
                textbox(slide, Inches(0.9), Inches(5.0), Inches(11.5), Inches(0.9), s["subtitle"], 18, GOLD,
                        font="Calibri")
            textbox(slide, Inches(0.9), Inches(6.5), Inches(11.5), Inches(0.5),
                    f"{meta.get('business', '')} · {meta.get('date', '')}", 12, MUTED, font="Calibri")
        elif s["kind"] == "end":
            if meta.get("signature") and brand.LOGO.exists():
                slide.shapes.add_picture(str(brand.LOGO), Inches(5.6), Inches(1.2), height=Inches(2.6))
            textbox(slide, Inches(0.5), Inches(4.4), Inches(12.3), Inches(1), s["title"], 40, CREAM, True,
                    align=PP_ALIGN.CENTER)
            textbox(slide, Inches(0.5), Inches(5.4), Inches(12.3), Inches(0.6), meta.get("business", ""), 16, GOLD,
                    font="Calibri", align=PP_ALIGN.CENTER)
        else:
            textbox(slide, Inches(0.8), Inches(0.5), Inches(11.2), Inches(1.1), s["title"], 30, CREAM, True)
            line = slide.shapes.add_shape(1, Inches(0.8), Inches(1.55), Inches(1.2), Emu(38000))
            line.fill.solid()
            line.fill.fore_color.rgb = GOLD
            line.line.fill.background()
            if s["kind"] == "table":
                rows = s["table"]
                ncol = max(len(r) for r in rows)
                shape = slide.shapes.add_table(len(rows), ncol, Inches(0.8), Inches(1.9), Inches(11.7),
                                               Inches(0.45) * len(rows))
                for ri, row in enumerate(rows):
                    for ci in range(ncol):
                        cell = shape.table.cell(ri, ci)
                        cell.text = tk.truncate_words(tk.strip_md(row[ci] if ci < len(row) else ""), 18)
                        para = cell.text_frame.paragraphs[0]
                        para.font.size = Pt(13 if ri == 0 else 12)
                        para.font.bold = ri == 0
                        para.font.color.rgb = INK if ri == 0 else CREAM
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = GOLD if ri == 0 else (
                            RGBColor(0x1A, 0x25, 0x40) if ri % 2 else RGBColor(0x24, 0x31, 0x4F))
            else:
                tb = slide.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(11.7), Inches(5.0))
                tf = tb.text_frame
                tf.word_wrap = True
                bullets = s.get("bullets") or []
                size = 22 if len(bullets) <= 4 else 19
                for bi, b in enumerate(bullets):
                    p = tf.paragraphs[0] if bi == 0 else tf.add_paragraph()
                    p.text = "◆  " + b
                    p.font.size = Pt(size)
                    p.font.name = "Calibri"
                    p.font.color.rgb = CREAM
                    p.space_after = Pt(14)
            textbox(slide, Inches(0.8), Inches(6.9), Inches(4), Inches(0.4), f"{i:02d} / {total:02d}", 11, MUTED,
                    font="Calibri")
            if meta.get("signature") and brand.LOGO_MARK.exists():
                slide.shapes.add_picture(str(brand.LOGO_MARK), Inches(12.1), Inches(6.75), height=Inches(0.5))
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"][:4000]
    prs.core_properties.title = doc.get("title", "")
    prs.save(str(path))
    return total


@plugin("slide_deck", "Slide Deck (PPT)", "documents", "render",
        "Builds a branded 16:9 PowerPoint deck with agenda, content and table slides and speaker notes.",
        icon="slides", keywords=["ppt", "pptx", "powerpoint", "slides", "presentation", "deck", "pitch deck"],
        params={"slides": "int"}, example="Make a 12-slide presentation on digital marketing for car dealers")
def slide_deck(ctx, slides: int = 0):
    doc = _need_doc(ctx)
    path = ctx.out_path(_fname(doc, ".pptx"))
    n = render_pptx(doc, path, _meta(ctx), int(slides or 0))
    ctx.add_output(path, f"Slide deck · {n} slides")
    ctx.log(f"Presentation created with {n} slides.")


# ====================================================================== Excel
TEMPLATES = {
    "budget": ["budget", "monthly budget", "household budget"],
    "expense": ["expense", "expenses", "spending"],
    "inventory": ["inventory", "stock"],
    "sales": ["sales tracker", "sales", "orders"],
    "attendance": ["attendance"],
    "project": ["project tracker", "task tracker", "to-do", "todo", "tasks", "project plan"],
    "crm": ["lead", "leads", "crm", "customer list", "clients list"],
}


def _xl_style(ws, header_row=1, widths=None, currency_cols=(), fmt='"₹"#,##0.00'):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    gold = PatternFill("solid", fgColor="1A2540")
    thin = Side(style="thin", color="D9CDB2")
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="EAD7A1")
        cell.fill = gold
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    for col in currency_cols:
        for row in ws.iter_rows(min_row=header_row + 1, min_col=col, max_col=col):
            for c in row:
                c.number_format = fmt


def _currency_fmt(ctx):
    country = (ctx.profile.get("country") or "India").lower()
    if "india" in country:
        return '"₹"#,##0.00'
    if country in ("usa", "united states", "us"):
        return '"$"#,##0.00'
    if "kingdom" in country or country == "uk":
        return '"£"#,##0.00'
    return "#,##0.00"


def _excel_template(wb, kind, ctx):
    from openpyxl.worksheet.datavalidation import DataValidation
    fmt = _currency_fmt(ctx)
    ws = wb.active
    today = dt.date.today()
    if kind == "budget":
        ws.title = "Budget"
        ws.append(["Category", "Type", "Planned", "Actual", "Difference", "Notes"])
        rows = [("Salary / Revenue", "Income"), ("Other income", "Income"), ("Rent", "Expense"),
                ("Utilities", "Expense"), ("Groceries / Supplies", "Expense"), ("Transport / Fuel", "Expense"),
                ("Marketing", "Expense"), ("Loan EMI", "Expense"), ("Insurance", "Expense"), ("Savings", "Expense"),
                ("Miscellaneous", "Expense")]
        for i, (c, t) in enumerate(rows, start=2):
            ws.append([c, t, 0, 0, f"=C{i}-D{i}", ""])
        n = len(rows) + 1
        ws.append([])
        ws.append(["Total income (actual)", "", "", f'=SUMIF(B2:B{n},"Income",D2:D{n})'])
        ws.append(["Total expenses (actual)", "", "", f'=SUMIF(B2:B{n},"Expense",D2:D{n})'])
        ws.append(["Net", "", "", f"=D{n + 2}-D{n + 3}"])
        _xl_style(ws, widths=[28, 12, 14, 14, 14, 30], currency_cols=(3, 4, 5), fmt=fmt)
        for r in range(n + 2, n + 5):
            ws.cell(row=r, column=4).number_format = fmt
    elif kind == "expense":
        ws.title = "Expenses"
        ws.append(["Date", "Category", "Description", "Amount", "Payment mode", "Receipt?"])
        for i in range(2, 102):
            ws.cell(row=i, column=1).number_format = "DD-MM-YYYY"
        dv = DataValidation(type="list", formula1='"Cash,UPI,Card,Bank transfer,Cheque"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("E2:E500")
        _xl_style(ws, widths=[14, 18, 36, 14, 16, 10], currency_cols=(4,), fmt=fmt)
        s = wb.create_sheet("Summary")
        s.append(["Category", "Total"])
        for i, cat in enumerate(["Rent", "Utilities", "Supplies", "Travel", "Food", "Marketing", "Salaries", "Other"], 2):
            s.append([cat, f'=SUMIF(Expenses!B:B,A{i},Expenses!D:D)'])
        s.append(["Grand total", "=SUM(B2:B9)"])
        _xl_style(s, widths=[20, 16], currency_cols=(2,), fmt=fmt)
    elif kind == "inventory":
        ws.title = "Inventory"
        ws.append(["SKU", "Item", "Category", "Qty in stock", "Reorder level", "Unit cost", "Stock value", "Status"])
        for i in range(2, 52):
            ws.cell(row=i, column=7, value=f'=IF(D{i}="","",D{i}*F{i})')
            ws.cell(row=i, column=8, value=f'=IF(D{i}="","",IF(D{i}<=E{i},"Reorder","OK"))')
        _xl_style(ws, widths=[12, 30, 16, 12, 12, 12, 14, 10], currency_cols=(6, 7), fmt=fmt)
    elif kind == "sales":
        ws.title = "Sales"
        ws.append(["Date", "Invoice #", "Customer", "Product / Service", "Qty", "Unit price", "Total", "Status"])
        for i in range(2, 202):
            ws.cell(row=i, column=7, value=f'=IF(E{i}="","",E{i}*F{i})')
        dv = DataValidation(type="list", formula1='"Paid,Pending,Overdue,Cancelled"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("H2:H500")
        _xl_style(ws, widths=[12, 12, 24, 30, 8, 12, 14, 12], currency_cols=(6, 7), fmt=fmt)
    elif kind == "attendance":
        ws.title = f"Attendance {today.strftime('%b %Y')}"
        ws.append(["Name"] + [str(d) for d in range(1, 32)] + ["Present", "Absent"])
        for i in range(2, 32):
            ws.cell(row=i, column=33, value=f'=COUNTIF(B{i}:AF{i},"P")')
            ws.cell(row=i, column=34, value=f'=COUNTIF(B{i}:AF{i},"A")')
        dv = DataValidation(type="list", formula1='"P,A,L,H"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("B2:AF31")
        _xl_style(ws, widths=[24] + [4] * 31 + [9, 9])
    elif kind == "project":
        ws.title = "Tasks"
        ws.append(["Task", "Owner", "Start", "Due", "Priority", "Status", "% Done", "Notes"])
        for i in range(2, 102):
            ws.cell(row=i, column=7).number_format = "0%"
        for rng, opts in (("E2:E500", '"High,Medium,Low"'), ("F2:F500", '"Not started,In progress,Blocked,Done"')):
            dv = DataValidation(type="list", formula1=opts, allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(rng)
        _xl_style(ws, widths=[36, 16, 12, 12, 10, 14, 9, 30])
    else:  # crm
        ws.title = "Leads"
        ws.append(["Name", "Phone", "Email", "Source", "Stage", "Value", "Next follow-up", "Notes"])
        dv = DataValidation(type="list", formula1='"New,Contacted,Qualified,Proposal,Won,Lost"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("E2:E1000")
        _xl_style(ws, widths=[22, 15, 26, 14, 12, 12, 15, 30], currency_cols=(6,), fmt=fmt)


def write_sheets(path, sheets: list, ctx=None):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sh in sheets:
        ws = wb.create_sheet(re.sub(r"[\\/*?:\[\]]", "", str(sh.get("name") or "Sheet"))[:31] or "Sheet")
        header = [str(h) for h in sh.get("header", [])]
        if header:
            ws.append(header)
        for row in sh.get("rows", []):
            ws.append([_num(v) for v in row])
        widths = []
        for ci in range(1, (len(header) or 1) + 1):
            col_vals = [len(str(ws.cell(row=r, column=ci).value or "")) for r in range(1, min(ws.max_row, 200) + 1)]
            widths.append(min(60, max(10, max(col_vals or [10]) + 2)))
        if header:
            _xl_style(ws, widths=widths)
    wb.save(str(path))


def _num(v):
    if isinstance(v, str):
        s = v.strip().replace(",", "")
        if re.fullmatch(r"-?\d+", s) and len(s) < 15 and not (s.startswith("0") and len(s) > 1):
            return int(s)
        if re.fullmatch(r"-?\d+\.\d+", s):
            return float(s)
    return v


@plugin("excel_creator", "Excel Builder", "documents", "render",
        "Creates spreadsheets: budgets, expense trackers, inventory, sales, attendance, task trackers, CRMs — with "
        "formulas, dropdowns and formatting — or exports tables from generated content.",
        icon="table", keywords=["excel", "xlsx", "spreadsheet", "sheet", "tracker", "budget", "inventory", "attendance"],
        params={"template": "budget|expense|inventory|sales|attendance|project|crm"},
        example="Create an inventory tracker spreadsheet for my spare parts shop")
def excel_creator(ctx, template: str = ""):
    import openpyxl
    low = ctx.problem.lower()
    template = template or next((k for k, kws in TEMPLATES.items() if any(w in low for w in kws)), "")
    sheets = ctx.state.get("sheets")
    name = "spreadsheet"
    if sheets:
        path = ctx.out_path(tk.slugify(ctx.state.get("sheets_name", "data")) + ".xlsx")
        write_sheets(path, sheets, ctx)
    elif template:
        wb = openpyxl.Workbook()
        _excel_template(wb, template, ctx)
        path = ctx.out_path(f"{template}-tracker.xlsx")
        wb.save(str(path))
        name = f"{template.title()} spreadsheet (with formulas)"
    elif ctx.state.get("doc") and any(b["type"] == "table" for s in ctx.state["doc"]["sections"] for b in s["blocks"]):
        sheets = []
        for s in ctx.state["doc"]["sections"]:
            for b in s["blocks"]:
                if b["type"] == "table" and b["rows"]:
                    sheets.append({"name": s["heading"][:28] or f"Table {len(sheets) + 1}",
                                   "header": b["rows"][0], "rows": b["rows"][1:]})
        path = ctx.out_path(_fname(ctx.state["doc"], ".xlsx"))
        write_sheets(path, sheets, ctx)
    elif ctx.llm.available():
        ctx.log("Designing a custom spreadsheet with the AI engine…")
        data = ctx.llm.complete_json(
            f"User request: {ctx.problem}\nDesign a useful spreadsheet. Return JSON "
            '{"name": str, "sheets": [{"name": str, "header": [str], "rows": [[str|number]]}]} with realistic '
            "example rows (max 30 per sheet). Do not invent facts about real people.", max_tokens=4000)
        sheets = [s for s in (data.get("sheets") or []) if isinstance(s, dict)][:6] if isinstance(data, dict) else []
        if not sheets:
            raise PluginError("Could not design the spreadsheet — try naming a template like budget or inventory.")
        path = ctx.out_path(tk.slugify(data.get("name", "spreadsheet")) + ".xlsx")
        write_sheets(path, sheets, ctx)
    else:
        wb = openpyxl.Workbook()
        _excel_template(wb, "project", ctx)
        path = ctx.out_path("tracker.xlsx")
        wb.save(str(path))
        ctx.log("No specific template recognised — created a task tracker. Mention budget, expense, inventory, sales, "
                "attendance or leads for other templates.", "warn")
    ctx.add_output(path, name if name != "spreadsheet" else "Spreadsheet")
    ctx.log("Spreadsheet created.")


# ====================================================================== Website
def _site_content(ctx):
    p = ctx.profile
    key, ind = industries.detect(" ".join([p.get("industry", ""), p.get("description", ""), ctx.problem]))
    content = {"name": ctx.business, "tagline": p.get("tagline") or ind["tagline"],
               "about": p.get("description") or f"{ctx.business} is a trusted name in {ind['label'].lower()}. "
                                                 "We combine expertise, honest pricing and genuine care for every customer.",
               "services": [{"title": t, "text": d} for t, d in ind["services"]],
               "why": ["Experienced, friendly team", "Transparent pricing — no surprises",
                       "Quality checked at every step", "Quick response on call and WhatsApp"],
               "cta": "Get in touch", "color": p.get("brand_color") or ind["color"]}
    if ctx.llm.available():
        try:
            data = ctx.llm.complete_json(
                f"Write website copy for this business. {ctx.problem}\nContext: {p}\nReturn JSON "
                '{"tagline": str, "about": str (60-90 words), "services": [{"title": str, "text": str}] (4-6), '
                '"why": [str] (4), "cta": str}. No testimonials, no invented awards, statistics or claims.',
                max_tokens=1800)
            if isinstance(data, dict):
                for k in ("tagline", "about", "cta"):
                    if isinstance(data.get(k), str) and data[k].strip():
                        content[k] = data[k].strip()
                if isinstance(data.get("services"), list) and data["services"]:
                    content["services"] = [{"title": str(s.get("title", ""))[:80], "text": str(s.get("text", ""))[:300]}
                                           for s in data["services"] if isinstance(s, dict)][:6]
                if isinstance(data.get("why"), list) and data["why"]:
                    content["why"] = [str(w)[:120] for w in data["why"]][:6]
        except LLMError as exc:
            ctx.log(f"AI copy unavailable ({exc}); using industry copy.", "warn")
    return content


@plugin("website_builder", "Website Builder", "documents", "render",
        "Generates a fast, responsive one-page website / landing page for your business as a single HTML file.",
        icon="globe", keywords=["website", "landing page", "web page", "webpage", "site", "portfolio site"],
        ai_enhanced=True, example="Build a landing page for my car detailing studio")
def website_builder(ctx):
    c = _site_content(ctx)
    p = ctx.profile
    e = lambda s: html.escape(str(s or ""))
    accent = c["color"] if re.fullmatch(r"#[0-9a-fA-F]{6}", c.get("color", "")) else "#C9A55C"
    contact = []
    if p.get("phone"):
        contact.append(f'<a href="tel:{e(re.sub(r"[^0-9+]", "", p["phone"]))}">{e(p["phone"])}</a>')
    if p.get("public_email"):
        contact.append(f'<a href="mailto:{e(p["public_email"])}">{e(p["public_email"])}</a>')
    if p.get("address"):
        contact.append(f"<span>{e(p['address'])}</span>")
    if p.get("instagram"):
        handle = p["instagram"].lstrip("@")
        contact.append(f'<a href="https://instagram.com/{e(handle)}" rel="noopener">@{e(handle)}</a>')
    if not contact:
        contact.append("<span>Add your phone, e-mail and address in your InEra profile.</span>")
    services = "".join(f'<article class="card"><h3>{e(s["title"])}</h3><p>{e(s["text"])}</p></article>'
                       for s in c["services"])
    why = "".join(f"<li>{e(w)}</li>" for w in c["why"])
    year = dt.date.today().year
    page = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{e(c['name'])} — {e(c['tagline'])}</title>
<meta name="description" content="{e(c['about'][:155])}">
<style>
:root{{--accent:{accent};--ink:#121110;--paper:#faf8f4;--muted:#6b665d}}
*{{box-sizing:border-box}}body{{margin:0;font-family:system-ui,-apple-system,"Segoe UI",Roboto,sans-serif;color:var(--ink);background:var(--paper);line-height:1.6}}
a{{color:inherit}}.wrap{{max-width:1100px;margin:0 auto;padding:0 20px}}
header{{position:sticky;top:0;background:rgba(250,248,244,.92);backdrop-filter:blur(8px);border-bottom:1px solid #e8e2d6;z-index:2}}
header .wrap{{display:flex;align-items:center;justify-content:space-between;height:64px}}
.brand{{font-family:Georgia,serif;font-size:1.35rem;font-weight:700;letter-spacing:.02em}}
nav a{{margin-left:22px;text-decoration:none;color:var(--muted);font-size:.95rem}}nav a:hover{{color:var(--ink)}}
.hero{{padding:110px 0 90px;background:linear-gradient(135deg,#121110,#2a251d);color:#f5efe2}}
.hero h1{{font-family:Georgia,serif;font-size:clamp(2.2rem,5vw,3.8rem);line-height:1.1;margin:0 0 18px;max-width:780px}}
.hero p{{font-size:1.15rem;color:#d9d0bd;max-width:620px;margin:0 0 34px}}
.btn{{display:inline-block;background:var(--accent);color:#121110;padding:14px 28px;border-radius:999px;font-weight:600;text-decoration:none}}
section{{padding:80px 0}}h2{{font-family:Georgia,serif;font-size:2rem;margin:0 0 12px}}.lead{{color:var(--muted);max-width:640px}}
.grid{{display:grid;grid-template-columns:repeat(auto-fit,minmax(240px,1fr));gap:20px;margin-top:34px}}
.card{{background:#fff;border:1px solid #ebe5d8;border-radius:16px;padding:26px;box-shadow:0 1px 2px rgba(0,0,0,.03)}}
.card h3{{margin:0 0 8px;font-size:1.1rem}}.card p{{margin:0;color:var(--muted)}}
.card::before{{content:"";display:block;width:36px;height:3px;background:var(--accent);margin-bottom:16px;border-radius:2px}}
.why{{background:#fff;border-top:1px solid #ebe5d8;border-bottom:1px solid #ebe5d8}}
.why ul{{list-style:none;padding:0;display:grid;grid-template-columns:repeat(auto-fit,minmax(230px,1fr));gap:14px;margin-top:28px}}
.why li{{padding:16px 18px;border-left:3px solid var(--accent);background:var(--paper);border-radius:8px}}
.contact{{text-align:center}}.contact .links{{display:flex;flex-wrap:wrap;gap:14px;justify-content:center;margin:26px 0}}
.contact .links>*{{padding:10px 18px;border:1px solid #e0d8c8;border-radius:999px;text-decoration:none;background:#fff}}
footer{{padding:28px 0;color:var(--muted);font-size:.9rem;text-align:center;border-top:1px solid #e8e2d6}}
@media (max-width:640px){{nav{{display:none}}.hero{{padding:80px 0 60px}}section{{padding:56px 0}}}}
</style></head>
<body>
<header><div class="wrap"><div class="brand">{e(c['name'])}</div>
<nav><a href="#services">Services</a><a href="#why">Why us</a><a href="#contact">Contact</a></nav></div></header>
<section class="hero"><div class="wrap"><h1>{e(c['tagline'])}</h1><p>{e(c['about'])}</p>
<a class="btn" href="#contact">{e(c['cta'])}</a></div></section>
<section id="services"><div class="wrap"><h2>What we do</h2><p class="lead">Services designed around what you actually need.</p>
<div class="grid">{services}</div></div></section>
<section id="why" class="why"><div class="wrap"><h2>Why choose {e(c['name'])}</h2><ul>{why}</ul></div></section>
<section id="contact" class="contact"><div class="wrap"><h2>Let's talk</h2>
<p class="lead" style="margin:0 auto">Reach out — we usually reply within a few hours.</p>
<div class="links">{''.join(contact)}</div><a class="btn" href="#contact">{e(c['cta'])}</a></div></section>
<footer>© {year} {e(c['name'])}. All rights reserved.</footer>
</body></html>"""
    path = ctx.out_path(tk.slugify(c["name"]) + "-website.html")
    path.write_text(page, encoding="utf-8")
    ctx.add_output(path, "Website (single HTML file — upload to any host)", "html")
    ctx.log("Website built. Upload the HTML file to any static host (or your cPanel) to go live.")


# ====================================================================== Invoice
def _words_indian(n: int) -> str:
    ones = ("Zero One Two Three Four Five Six Seven Eight Nine Ten Eleven Twelve Thirteen Fourteen Fifteen Sixteen "
            "Seventeen Eighteen Nineteen").split()
    tens = "Zero Ten Twenty Thirty Forty Fifty Sixty Seventy Eighty Ninety".split()

    def two(x):
        return ones[x] if x < 20 else tens[x // 10] + ("" if x % 10 == 0 else " " + ones[x % 10])

    def three(x):
        h, r = divmod(x, 100)
        return ((ones[h] + " Hundred" + (" " if r else "")) if h else "") + (two(r) if r else "")
    if n == 0:
        return "Zero"
    parts = []
    for div, name in ((10 ** 7, "Crore"), (10 ** 5, "Lakh"), (1000, "Thousand")):
        q, n = divmod(n, div)
        if q:
            parts.append((three(q) if q < 1000 else _words_indian(q)) + " " + name)
    if n:
        parts.append(three(n))
    return " ".join(parts)


def _parse_invoice(text: str) -> dict:
    data = {"client": "", "items": [], "tax_rate": 0.0}
    m = re.search(r"\b(?:to|for|client|bill to|customer)[:\s]+([A-Z][\w&.' ]{2,50}?)(?=[,.;:\n]| for | with | -|$)", text)
    if m:
        data["client"] = m.group(1).strip()
    dm = re.search(r"due (?:in|within|after) (\d{1,3}) days", text, re.I)
    if dm:
        data["due_days"] = int(dm.group(1))
    tm = re.search(r"(\d{1,2}(?:\.\d+)?)\s*%\s*(?:gst|tax|vat|igst)", text, re.I)
    if tm:
        data["tax_rate"] = float(tm.group(1))
    qty_first = re.compile(r"(?P<qty>\d{1,5})\s*(?:x|×|\*|nos?\.?|pcs?\.?|units?)?\s+(?P<desc>[A-Za-z][A-Za-z0-9 &/()'\-]{1,60}?)"
                           r"\s*(?:@|\bat\b|\bfor\b|-|–|:)\s*(?:₹|rs\.?|inr|\$)?\s*(?P<rate>\d[\d,]*(?:\.\d{1,2})?)(?!\s*%)", re.I)
    desc_first = re.compile(r"(?P<desc>[A-Za-z][A-Za-z0-9 &/()'\-]{2,60}?)\s*[-:–]?\s*(?:(?P<qty>\d{1,5})\s*[x×*]\s*)?"
                            r"(?:@\s*)?(?:₹|rs\.?|inr|\$)?\s*(?P<rate>\d[\d,]*(?:\.\d{1,2})?)(?!\s*%)", re.I)
    segments = re.split(r"[\n;]|,\s+|,(?=[A-Za-z₹])|\s+and\s+(?=[A-Za-z0-9])|\.\s+(?=[A-Za-z])", text)
    for seg in segments:
        seg = seg.rsplit(":", 1)[-1] if re.search(r"invoice|bill|estimate|quotation", seg.split(":")[0], re.I) else seg
        mm = qty_first.search(seg) or desc_first.search(seg)
        if not mm:
            continue
        desc = re.sub(r"^(?:and|for|with|item|items|invoice|create|make|generate|an?|the)\s+", "", mm.group("desc").strip(),
                      flags=re.I).strip(" -:")
        if re.search(r"\b(gst|tax|invoice|days|due|bill)\b", desc, re.I) or len(desc) < 3:
            continue
        rate = float(mm.group("rate").replace(",", ""))
        if rate <= 0:
            continue
        if data["client"] and desc.lower().startswith(data["client"].lower()):
            continue
        data["items"].append({"desc": desc[:1].upper() + desc[1:], "qty": int(mm.group("qty") or 1), "rate": rate})
    return data


@plugin("invoice_generator", "Invoice Generator", "data", "render",
        "Creates a professional invoice PDF with your business details, tax, totals and amount in words.",
        icon="receipt", keywords=["invoice", "bill", "receipt", "quotation", "estimate"],
        example="Invoice for Ravi Motors: brake pads 2 x 1500, wheel alignment 800, 18% GST")
def invoice_generator(ctx):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    data = _parse_invoice(ctx.problem)
    if ctx.llm.available():
        try:
            ai = ctx.llm.complete_json(
                f"Extract invoice data from: {ctx.problem}\nReturn JSON {{\"client\": str, \"client_address\": str, "
                '"items": [{"desc": str, "qty": number, "rate": number}], "tax_rate": number, "notes": str, '
                '"due_days": number}. Use only what the text states.', max_tokens=1200)
            if isinstance(ai, dict) and ai.get("items"):
                items = []
                for it in ai["items"]:
                    try:
                        items.append({"desc": str(it["desc"])[:120], "qty": float(it.get("qty") or 1),
                                      "rate": float(it.get("rate") or 0)})
                    except (KeyError, TypeError, ValueError):
                        continue
                if items:
                    data.update({k: ai.get(k) for k in ("client", "client_address", "notes", "due_days") if ai.get(k)})
                    data["items"] = items
                    data["tax_rate"] = float(ai.get("tax_rate") or data["tax_rate"] or 0)
        except (LLMError, ValueError):
            pass
    if not data["items"]:
        raise PluginError("I couldn't find line items. Write them like: “brake pads 2 x 1500, wheel alignment 800, 18% GST”.")
    p = ctx.profile
    f = brand.reportlab_fonts()
    india = "india" in (p.get("country") or "India").lower()
    cur = "₹" if india else (p.get("currency_symbol") or "")
    if not f["unicode"] and cur == "₹":
        cur = "Rs. "
    money = lambda v: f"{cur}{v:,.2f}"
    number = "INV-" + dt.date.today().strftime("%Y%m%d") + "-" + ctx.job_id[-4:].upper()
    sub = sum(i["qty"] * i["rate"] for i in data["items"])
    tax = round(sub * data["tax_rate"] / 100, 2)
    total = round(sub + tax, 2)
    gold, ink = colors.HexColor("#C9A55C"), colors.HexColor("#1A1814")
    S = lambda size=10, bold=False, color=ink, align=0: ParagraphStyle(
        "x", fontName=f["bold"] if bold else f["regular"], fontSize=size, leading=size * 1.35, textColor=color,
        alignment=align)
    E = lambda s: html.escape(brand.pdf_safe(str(s or "")))
    seller = [f"<b>{E(ctx.business)}</b>"] + [E(p[k]) for k in ("address", "phone", "public_email") if p.get(k)]
    if p.get("gstin"):
        seller.append("GSTIN: " + E(p["gstin"]))
    due = dt.date.today() + dt.timedelta(days=int(data.get("due_days") or 15))
    story = [Table([[Paragraph("<br/>".join(seller), S(10)),
                     Paragraph(f"<b>INVOICE</b><br/>{number}<br/>Date: {dt.date.today():%d %b %Y}<br/>"
                               f"Due: {due:%d %b %Y}", S(11, align=2))]], colWidths=[100 * mm, 70 * mm]),
             Spacer(1, 8 * mm),
             Paragraph("<b>Bill to</b><br/>" + E(data.get("client") or "Customer") +
                       ("<br/>" + E(data.get("client_address")) if data.get("client_address") else ""), S(10)),
             Spacer(1, 6 * mm)]
    rows = [["#", "Description", "Qty", "Rate", "Amount"]]
    for n, it in enumerate(data["items"], 1):
        q = it["qty"]
        rows.append([str(n), Paragraph(E(it["desc"]), S(9.5)), f"{q:g}", money(it["rate"]), money(q * it["rate"])])
    rows.append(["", "", "", "Subtotal", money(sub)])
    if data["tax_rate"]:
        label = f"GST {data['tax_rate']:g}%" if india else f"Tax {data['tax_rate']:g}%"
        rows.append(["", "", "", label, money(tax)])
    rows.append(["", "", "", "Total", money(total)])
    tbl = Table(rows, colWidths=[10 * mm, 82 * mm, 16 * mm, 30 * mm, 32 * mm])
    n_items = len(data["items"])
    tbl.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), f["regular"]), ("FONTSIZE", (0, 0), (-1, -1), 9.5),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1A2540")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), f["bold"]),
        ("ALIGN", (2, 0), (-1, -1), "RIGHT"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LINEBELOW", (0, 1), (-1, n_items), 0.4, colors.HexColor("#E2D8C3")),
        ("FONTNAME", (3, -1), (-1, -1), f["bold"]), ("LINEABOVE", (3, -1), (-1, -1), 1.2, gold),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    story += [tbl, Spacer(1, 5 * mm)]
    if india:
        rupees = int(total)
        paise = int(round((total - rupees) * 100))
        words = _words_indian(rupees) + " Rupees" + (f" and {_words_indian(paise)} Paise" if paise else "") + " Only"
        story.append(Paragraph(f"<b>Amount in words:</b> {E(words)}", S(9.5)))
        if data["tax_rate"]:
            story.append(Paragraph(f"For intra-state supply this GST splits equally into CGST and SGST "
                                   f"({data['tax_rate'] / 2:g}% each); use IGST for inter-state supply.", S(8.5, color=colors.HexColor("#7A7263"))))
    if p.get("upi_id") or p.get("bank_details"):
        story += [Spacer(1, 5 * mm), Paragraph("<b>Payment details</b><br/>" +
                                               (f"UPI: {E(p['upi_id'])}<br/>" if p.get("upi_id") else "") +
                                               E(p.get("bank_details", "")), S(9.5))]
    if data.get("notes"):
        story += [Spacer(1, 4 * mm), Paragraph("<b>Notes</b><br/>" + E(data["notes"]), S(9.5))]
    story += [Spacer(1, 12 * mm), Paragraph("Thank you for your business.", S(10, color=colors.HexColor("#8A6A32")))]
    path = ctx.out_path(f"{number}.pdf")
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm,
                      bottomMargin=18 * mm, title=number).build(story)
    ctx.add_output(path, f"Invoice {number} · {money(total)}")
    ctx.log(f"Invoice {number} created: {len(data['items'])} item(s), total {money(total)}. "
            "Check the details before sending — this is not tax advice.")


# ====================================================================== Data analysis
def _to_float(v):
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v) if not (isinstance(v, float) and math.isnan(v)) else None
    s = str(v).strip().replace(",", "").replace("₹", "").replace("$", "").replace("%", "")
    try:
        return float(s) if s else None
    except ValueError:
        return None


def _to_date(v):
    if isinstance(v, (dt.date, dt.datetime)):
        return v if isinstance(v, dt.datetime) else dt.datetime(v.year, v.month, v.day)
    s = str(v or "").strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y", "%Y/%m/%d", "%d %b %Y", "%b %Y", "%Y-%m"):
        try:
            return dt.datetime.strptime(s[:11].strip(), fmt)
        except ValueError:
            continue
    return None


def _chart(fig_fn, path):
    """Thread-safe chart rendering (no pyplot global state)."""
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.figure import Figure
    fig = Figure(figsize=(8, 4.4), dpi=130)
    FigureCanvasAgg(fig)
    ax = fig.add_subplot()
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        ax.spines[side].set_color("#8A8170")
    ax.tick_params(colors="#5A544A", labelsize=9)
    fig_fn(ax)
    fig.tight_layout()
    fig.savefig(str(path))


PALETTE = ["#C9A55C", "#24314F", "#8A6A32", "#5B8C85", "#B85C5C", "#6B7FD7", "#D9A441", "#7A7263"]


@plugin("data_analyst", "Data Analyst", "data", "content",
        "Analyses CSV/Excel files: column profiles, statistics, trends, top categories and charts, delivered as a "
        "report with a summary workbook.",
        icon="chart", keywords=["analyze", "analyse", "analysis", "csv", "dataset", "data", "statistics", "insights",
                                "trend", "dashboard"],
        needs=["files:csv,xlsx"], example="Analyze my sales.csv and tell me the trends")
def data_analyst(ctx):
    files = ctx.input_files([".csv", ".tsv", ".xlsx", ".xlsm"])
    if not files:
        raise PluginError("Upload a CSV or Excel file to analyse.")
    f = files[0]
    header, rows = tk.read_table_file(f["path"])
    if not rows:
        raise PluginError("The file has no data rows.")
    ncols = len(header)
    rows = [list(r) + [None] * (ncols - len(r)) for r in rows]
    cols = []
    for ci, name in enumerate(header):
        vals = [r[ci] for r in rows]
        nonempty = [v for v in vals if v not in (None, "")]
        nums = [x for x in (_to_float(v) for v in nonempty) if x is not None]
        dates = [x for x in (_to_date(v) for v in nonempty[:200]) if x is not None]
        if nonempty and len(dates) >= 0.8 * min(len(nonempty), 200) and len(nums) < 0.8 * len(nonempty):
            kind = "date"
        elif nonempty and len(nums) >= 0.8 * len(nonempty):
            kind = "number"
        else:
            kind = "text"
        cols.append({"name": name, "kind": kind, "missing": len(vals) - len(nonempty), "values": vals, "nums": nums,
                     "unique": len(set(map(str, nonempty)))})
    doc = tk.new_doc(f"Data Analysis: {f['name']}", f"{len(rows):,} rows × {ncols} columns")
    profile_rows = [["Column", "Type", "Missing", "Unique", "Summary"]]
    for c in cols:
        if c["kind"] == "number" and c["nums"]:
            summ = f"min {min(c['nums']):,.2f} · mean {statistics.fmean(c['nums']):,.2f} · max {max(c['nums']):,.2f}"
        elif c["kind"] == "text":
            top = Counter(str(v) for v in c["values"] if v not in (None, "")).most_common(1)
            summ = f"most common: {top[0][0][:30]} ({top[0][1]})" if top else ""
        else:
            summ = "dates"
        profile_rows.append([c["name"], c["kind"], str(c["missing"]), str(c["unique"]), summ])
    tk.add_section(doc, "Column Profile", table=profile_rows)
    insights = []
    numeric = [c for c in cols if c["kind"] == "number" and len(c["nums"]) > 1]
    cats = [c for c in cols if c["kind"] == "text" and 1 < c["unique"] <= max(30, len(rows) // 3)]
    dates = [c for c in cols if c["kind"] == "date"]
    stats_rows = [["Metric", "Count", "Sum", "Mean", "Median", "Std dev", "Min", "Max"]]
    for c in numeric:
        n = c["nums"]
        stats_rows.append([c["name"], str(len(n)), f"{sum(n):,.2f}", f"{statistics.fmean(n):,.2f}",
                           f"{statistics.median(n):,.2f}", f"{statistics.pstdev(n):,.2f}", f"{min(n):,.2f}", f"{max(n):,.2f}"])
    if numeric:
        tk.add_section(doc, "Statistics", table=stats_rows)
    sheets = [{"name": "Profile", "header": profile_rows[0], "rows": profile_rows[1:]}]
    if numeric:
        sheets.append({"name": "Statistics", "header": stats_rows[0], "rows": stats_rows[1:]})
    charts = []
    main = max(numeric, key=lambda c: sum(c["nums"]), default=None)
    if cats and main:
        cat = cats[0]
        ci, mi = header.index(cat["name"]), header.index(main["name"])
        agg = Counter()
        for r in rows:
            v = _to_float(r[mi])
            if v is not None and r[ci] not in (None, ""):
                agg[str(r[ci])] += v
        top = agg.most_common(10)
        if top:
            path = ctx.out_path(f"chart-{tk.slugify(main['name'])}-by-{tk.slugify(cat['name'])}.png")

            def bar(ax):
                labels = [t[0][:18] for t in top][::-1]
                ax.barh(labels, [t[1] for t in top][::-1], color=PALETTE[0])
                ax.set_title(f"{main['name']} by {cat['name']} (top {len(top)})", loc="left", fontweight="bold")
            _chart(bar, path)
            charts.append((path, f"{main['name']} by {cat['name']}"))
            share = top[0][1] / max(sum(agg.values()), 1e-9) * 100
            insights.append(f"**{top[0][0]}** leads {main['name']} with {top[0][1]:,.2f} ({share:.1f}% of total).")
            if 3 < len(agg):
                insights.append(f"Top 3 {cat['name']} values account for "
                                f"{sum(t[1] for t in top[:3]) / max(sum(agg.values()), 1e-9) * 100:.1f}% of {main['name']}.")
            sheets.append({"name": f"By {cat['name']}"[:31], "header": [cat["name"], main["name"]],
                           "rows": [[k, round(v, 2)] for k, v in agg.most_common()]})
    if dates and main:
        di, mi = header.index(dates[0]["name"]), header.index(main["name"])
        series = Counter()
        for r in rows:
            d0, v = _to_date(r[di]), _to_float(r[mi])
            if d0 and v is not None:
                series[d0.strftime("%Y-%m")] += v
        if len(series) >= 2:
            keys = sorted(series)
            path = ctx.out_path(f"trend-{tk.slugify(main['name'])}.png")

            def line(ax):
                ax.plot(keys, [series[k] for k in keys], color=PALETTE[0], linewidth=2.4, marker="o")
                ax.set_title(f"{main['name']} per month", loc="left", fontweight="bold")
                ax.tick_params(axis="x", rotation=45)
            _chart(line, path)
            charts.append((path, f"{main['name']} per month"))
            first, last = series[keys[0]], series[keys[-1]]
            if first:
                insights.append(f"{main['name']} moved from {first:,.2f} in {keys[0]} to {last:,.2f} in {keys[-1]} "
                                f"({(last - first) / abs(first) * 100:+.1f}%).")
            best = max(keys, key=lambda k: series[k])
            insights.append(f"Best month for {main['name']}: **{best}** ({series[best]:,.2f}).")
            sheets.append({"name": "Monthly trend", "header": ["Month", main["name"]],
                           "rows": [[k, round(series[k], 2)] for k in keys]})
    for c in numeric[:3]:
        if c is main and charts:
            continue
        path = ctx.out_path(f"distribution-{tk.slugify(c['name'])}.png")
        _chart(lambda ax, c=c: (ax.hist(c["nums"], bins=min(20, max(5, len(c["nums"]) // 5)), color=PALETTE[2]),
                                ax.set_title(f"Distribution of {c['name']}", loc="left", fontweight="bold")), path)
        charts.append((path, f"Distribution of {c['name']}"))
    if len(numeric) >= 2:
        import numpy as np
        best_pair, best_r = None, 0
        for i in range(len(numeric)):
            for j in range(i + 1, len(numeric)):
                a_idx, b_idx = header.index(numeric[i]["name"]), header.index(numeric[j]["name"])
                pairs = [(_to_float(r[a_idx]), _to_float(r[b_idx])) for r in rows]
                pairs = [p for p in pairs if p[0] is not None and p[1] is not None]
                if len(pairs) > 5:
                    a, b = np.array(pairs).T
                    if a.std() and b.std():
                        r = float(np.corrcoef(a, b)[0, 1])
                        if abs(r) > abs(best_r):
                            best_pair, best_r = (numeric[i]["name"], numeric[j]["name"]), r
        if best_pair and abs(best_r) >= 0.3:
            strength = "strong" if abs(best_r) >= 0.7 else "moderate"
            insights.append(f"{best_pair[0]} and {best_pair[1]} show a {strength} "
                            f"{'positive' if best_r > 0 else 'negative'} correlation (r = {best_r:.2f}) — "
                            "correlation is not causation.")
    missing = [c for c in cols if c["missing"]]
    if missing:
        insights.append("Missing values in: " + ", ".join(f"{c['name']} ({c['missing']})" for c in missing[:6]) + ".")
    if ctx.llm.available() and insights:
        try:
            narrative = ctx.llm.complete(
                "You are a data analyst. Based only on these computed facts, write 4-6 bullet insights and 3 "
                "recommended actions for a business owner. Facts:\n" + "\n".join(insights) +
                "\n\nStatistics table:\n" + "\n".join(", ".join(r) for r in stats_rows[:12]), max_tokens=900)
            nd = tk.markdown_to_doc(narrative)
            blocks = [b for s in nd["sections"] for b in s["blocks"]]
            if blocks:
                doc["sections"].insert(0, {"heading": "Executive Summary", "blocks": blocks})
        except LLMError:
            pass
    if insights:
        tk.add_section(doc, "Key Insights", bullets=insights)
    if charts:
        sec = tk.add_section(doc, "Charts")
        for pth, cap in charts:
            sec["blocks"].append({"type": "image", "path": str(pth), "caption": cap})
            ctx.add_output(pth, cap)
    ctx.state["doc"] = doc
    ctx.state["sheets"] = sheets
    ctx.state["sheets_name"] = "analysis-" + Path(f["name"]).stem
    report = ctx.out_path(f"analysis-{tk.slugify(Path(f['name']).stem)}.pdf")
    render_pdf(doc, report, _meta(ctx))
    ctx.add_output(report, "Analysis report (PDF)")
    xl = ctx.out_path(f"analysis-{tk.slugify(Path(f['name']).stem)}.xlsx")
    write_sheets(xl, sheets, ctx)
    ctx.add_output(xl, "Analysis workbook")
    ctx.add_text_output("analysis.md", tk.doc_to_markdown(doc), "Analysis summary", "markdown")
    ctx.log(f"Analysed {len(rows):,} rows, {ncols} columns, produced {len(charts)} chart(s).")


# ====================================================================== Charts
@plugin("chart_maker", "Chart Maker", "data", "render",
        "Draws bar, line or pie charts from numbers in your message or the first two columns of an uploaded table.",
        icon="chart", keywords=["chart", "graph", "plot", "pie chart", "bar chart", "line chart", "visualize"],
        params={"type": "bar|line|pie"}, example="Make a bar chart: Jan 120, Feb 150, Mar 90, Apr 180")
def chart_maker(ctx, type: str = ""):
    low = ctx.problem.lower()
    ctype = type or ("pie" if "pie" in low else "line" if ("line" in low or "trend" in low) else "bar")
    pairs = []
    tables = ctx.input_files([".csv", ".tsv", ".xlsx"])
    if tables:
        header, rows = tk.read_table_file(tables[0]["path"])
        for r in rows[:50]:
            if len(r) >= 2 and _to_float(r[1]) is not None:
                pairs.append((str(r[0])[:24], _to_float(r[1])))
        title = f"{header[1]} by {header[0]}" if len(header) > 1 else "Chart"
    else:
        body = re.sub(r"^.*?:", "", ctx.problem, count=1) if ":" in ctx.problem else ctx.problem
        if ":" in ctx.problem:
            head = re.sub(r"\b(make|create|draw|generate|plot|give me|show|a|an|the|me|of|for|bar|line|pie|chart|graph)\b",
                          " ", ctx.problem.split(":")[0], flags=re.I)
            head = re.sub(r"\s+", " ", head).strip()
            if head:
                ctx.state["topic"] = head
        for m in re.finditer(r"([A-Za-z][A-Za-z .&'\-]{0,24}?)\s*[:=\-]?\s*(-?\d[\d,]*(?:\.\d+)?)(?!\s*%?\s*(?:x|\d))", body):
            label = m.group(1).strip(" -:=")
            if label.lower() in ("make a bar chart", "chart", "bar chart", "pie chart", "line chart") or not label:
                continue
            pairs.append((label[:24], float(m.group(2).replace(",", ""))))
        title = tk.title_case(ctx.topic() or "Chart")
    if len(pairs) < 2:
        raise PluginError("I need at least two label–number pairs, e.g. “Jan 120, Feb 150, Mar 90” or an uploaded CSV.")
    labels, values = [p[0] for p in pairs], [p[1] for p in pairs]
    path = ctx.out_path(f"{tk.slugify(title)}-{ctype}.png")

    def draw(ax):
        if ctype == "pie":
            ax.pie(values, labels=labels, autopct="%1.0f%%", colors=PALETTE * 4, startangle=90,
                   wedgeprops={"edgecolor": "white", "linewidth": 1.5})
            ax.axis("equal")
        elif ctype == "line":
            ax.plot(labels, values, color=PALETTE[0], linewidth=2.6, marker="o")
            ax.fill_between(range(len(values)), values, alpha=0.12, color=PALETTE[0])
        else:
            bars = ax.bar(labels, values, color=PALETTE[0])
            ax.bar_label(bars, fmt="%g", padding=3, fontsize=9)
        ax.set_title(title, loc="left", fontweight="bold")
        if ctype != "pie" and len(labels) > 6:
            ax.tick_params(axis="x", rotation=40)
    _chart(draw, path)
    ctx.add_output(path, f"{ctype.title()} chart")
    ctx.state.setdefault("charts", []).append(str(path))
