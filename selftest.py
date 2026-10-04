"""InEra self-test: runs the planner and every offline plugin end to end in a
throw-away data folder and prints PASS/FAIL. No network or AI engine needed.

    python selftest.py            # quick (skips video)
    python selftest.py --full     # includes the narrated video
"""
import sys
sys.dont_write_bytecode = True  # keep the app folder flat: no __pycache__

import csv
import json
import os
import shutil
import sys
import tempfile
import time
from pathlib import Path

TMP = Path(tempfile.mkdtemp(prefix="inera-selftest-"))
os.environ["INERA_DATA_DIR"] = str(TMP / "data")
os.environ["LLM_PROVIDER"] = "none"
os.environ["WEB_RESEARCH"] = "false"
os.environ["SECRET_KEY"] = "selftest-only"

import config  # noqa: E402
import db  # noqa: E402
import engine  # noqa: E402
import planner  # noqa: E402

SAMPLE = """Solar Energy

Solar energy is the radiant light and heat from the Sun that is harnessed using technologies such as solar panels.
A photovoltaic cell is a device that converts light directly into electricity using the photovoltaic effect.
In 1839 Edmond Becquerel discovered the photovoltaic effect while experimenting with electrolytic cells.
In 1954 Bell Laboratories demonstrated the first practical silicon solar cell.

How Panels Work

Solar panels contain many photovoltaic cells made of silicon. When sunlight strikes a cell, it frees electrons and creates
an electric current. An inverter is a device that converts direct current into alternating current for homes.
Net metering is a billing arrangement that credits owners for electricity they export to the grid.

Benefits and Limits

Solar power produces no emissions while operating and has low running costs. Output depends on sunlight, so
batteries or the grid are needed at night. In 2023 solar power was the fastest growing source of new electricity worldwide.
Panel efficiency for typical modules ranges between 18 and 23 percent.
"""


def make_samples(folder: Path) -> dict:
    from PIL import Image, ImageDraw
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    import docx
    folder.mkdir(parents=True, exist_ok=True)
    ss = getSampleStyleSheet()
    story = []
    for block in SAMPLE.strip().split("\n\n"):
        style = ss["Heading2"] if len(block.split()) < 5 else ss["BodyText"]
        story += [Paragraph(block.replace("\n", " "), style), Spacer(1, 6)]
    SimpleDocTemplate(str(folder / "solar.pdf"), pagesize=A4).build(story)
    SimpleDocTemplate(str(folder / "extra.pdf"), pagesize=A4).build([Paragraph("Appendix page", ss["Title"])])
    d = docx.Document()
    d.add_heading("Service policy", 0)
    d.add_paragraph("All repairs carry a six month warranty on parts and labour.")
    d.save(str(folder / "policy.docx"))
    for i, colour in enumerate([(30, 60, 120), (150, 40, 40)]):
        im = Image.new("RGB", (1200, 900), colour)
        ImageDraw.Draw(im).rectangle((300, 400, 900, 650), fill=(20, 20, 20))
        im.save(folder / f"photo{i + 1}.jpg", quality=88)
    with open(folder / "sales.csv", "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["Date", "Branch", "Units", "Revenue"])
        for i in range(60):
            w.writerow([f"2026-{1 + i % 6:02d}-{1 + i % 27:02d}", ["North", "South", "East", "West"][i % 4], 1 + i % 5,
                        (1 + i % 5) * 1200])
    return {p.name: p for p in folder.iterdir()}


def main():
    full = "--full" in sys.argv
    db.init()
    uid = "u_selftest"
    db.execute("INSERT INTO users (id,email,pw_hash,name,profile_json,created_at) VALUES (?,?,?,?,?,?)",
               (uid, "selftest@example.com", "x", "Self Test",
                db.jdump({"business_name": "Test Motors", "industry": "Automobile", "timezone": "Asia/Kolkata",
                          "country": "India", "city": "Belagavi"}), time.time()))
    user = db.get_user(uid)
    samples = make_samples(TMP / "samples")

    def upload(name):
        fid = db.new_id("f_")
        dest = config.FILES_DIR / uid / "uploads" / f"{fid}{Path(name).suffix}"
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(samples[name], dest)
        db.execute("INSERT INTO files (id,user_id,name,path,mime,size,origin,created_at) VALUES (?,?,?,?,?,?,?,?)",
                   (fid, uid, name, str(dest), "application/octet-stream", dest.stat().st_size, "upload", time.time()))
        return fid

    cases = [
        ("Make a study notebook from this PDF with a quiz, flashcards and a mind map", ["solar.pdf"],
         ["notebook_builder", "pdf_creator", "quiz_mcq", "flashcards", "mind_map"]),
        ("Turn this PDF into a 10 slide presentation, a Word file and a podcast", ["solar.pdf"],
         ["slide_deck", "word_creator", "audio_overview"]),
        ("Convert the attached PDF to Word", ["solar.pdf"], ["pdf_to_word"]),
        ("Convert this Word file to PDF", ["policy.docx"], ["office_to_pdf"]),
        ("Merge these PDFs, add page numbers and a DRAFT watermark", ["solar.pdf", "extra.pdf"],
         ["merge_pdf", "page_numbers", "watermark_pdf"]),
        ("Compress this PDF and convert it to JPG images", ["solar.pdf"], ["compress_pdf", "pdf_to_images"]),
        ("Protect this PDF with password Test@2026", ["extra.pdf"], ["protect_pdf"]),
        ("Make a PDF from these photos", ["photo1.jpg", "photo2.jpg"], ["images_to_pdf"]),
        ("Resize these photos for an Instagram story and convert to webp", ["photo1.jpg", "photo2.jpg"], ["image_tools"]),
        ("Analyze the attached sales CSV and show the trends", ["sales.csv"], ["data_analyst"]),
        ("Make a bar chart: Jan 120, Feb 150, Mar 90", [], ["chart_maker"]),
        ("Create an inventory tracker spreadsheet", [], ["excel_creator"]),
        ("Invoice for Ravi Motors: brake pads 2 x 1500, wheel alignment 800, 18% GST", [], ["invoice_generator"]),
        ("Build a landing page for my business", [], ["website_builder"]),
        ("Build a Python FastAPI app to manage bookings", [], ["code_project_builder"]),
        ("Make a QR code for https://example.com", [], ["qr_generator"]),
        ("Plan 30 days of Instagram content for my business", [], ["content_calendar"]),
        ("Create an Instagram post about tyre care", [], ["social_post_designer"]),
        ("Write a friendly email to my customers announcing our festive offer", [], ["content_writer"]),
        ("Summarize the attached PDF into key points", ["solar.pdf"], ["summarizer"]),
    ]
    if full:
        cases.append(("Make a video overview of this PDF", ["solar.pdf"], ["video_overview"]))
    passed = failed = 0
    for problem, files, must in cases:
        t0 = time.time()
        ids = [upload(f) for f in files]
        job_id = db.new_id("j_")
        clean, secrets = planner.split_secrets(problem)
        db.execute("INSERT INTO jobs (id,user_id,problem,status,input_ids_json,source,created_at) VALUES (?,?,?,?,?,?,?)",
                   (job_id, uid, clean, "queued", db.jdump(ids), "manual", time.time()))
        engine._pending[job_id] = problem
        if secrets:
            engine._secrets[job_id] = secrets
        try:
            engine.run_job(job_id)
        except Exception as exc:  # pragma: no cover
            print(f"FAIL  {problem[:60]} — crashed: {exc}")
            failed += 1
            continue
        job = db.query_one("SELECT * FROM jobs WHERE id=?", (job_id,))
        plan = json.loads(job["plan_json"])
        steps = {s["plugin"]: s for s in plan["steps"]}
        outputs = json.loads(job["outputs_json"])
        missing = [m for m in must if m not in steps]
        bad = [f"{k}: {v.get('error')}" for k, v in steps.items() if v["status"] == "failed"]
        files_ok = all(Path(config.FILES_DIR / uid / job_id / o["name"]).exists() for o in outputs)
        ok = not missing and not bad and outputs and files_ok
        passed += ok
        failed += not ok
        print(f"{'PASS' if ok else 'FAIL'}  {problem[:66]:<66} {len(outputs):>2} outputs  {time.time() - t0:5.1f}s")
        for line in missing:
            print(f"      planner missed: {line}")
        for line in bad:
            print(f"      {line}")
    sched = planner.parse_schedule("every day at 8 am", "Asia/Kolkata")
    sched_ok = sched == {"type": "daily", "times": ["08:00"]}
    print(f"{'PASS' if sched_ok else 'FAIL'}  schedule parser: every day at 8 am -> {sched}")
    passed += sched_ok
    failed += not sched_ok
    import brand
    print(f"\nEngines: speech={brand.tts_engine()} ffmpeg={'yes' if brand.ffmpeg_bin() else 'no'} "
          f"office={'yes' if __import__('plg_pdftools')._soffice() else 'no (built-in fallback)'}")
    print(f"Result: {passed} passed, {failed} failed")
    shutil.rmtree(TMP, ignore_errors=True)
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
