"""PDF & file tools (iLovePDF-style) — everything runs locally, nothing is
uploaded to third parties. License-clean libraries only (no AGPL)."""
import io
import os
import re
import shutil
import statistics
import subprocess
import tempfile
import zipfile
from functools import lru_cache
from pathlib import Path

from PIL import Image, ImageOps

import brand
import config
import textkit as tk
from plg_docs import _meta, render_pdf
from registry import PluginError, plugin

OFFICE_EXT = {".docx", ".doc", ".odt", ".rtf", ".pptx", ".ppt", ".odp", ".xlsx", ".xls", ".ods", ".txt", ".md",
              ".html", ".htm", ".csv"}


def _pdfs(ctx, minimum=1):
    """PDF tools chain: each tool works on the previous tool's result when there is one
    (merge -> number pages -> watermark -> protect), otherwise on the uploaded PDFs."""
    chain = [c for c in ctx.state.get("pdf_chain", []) if Path(c["path"]).exists()]
    files = chain if (chain and minimum == 1) else (ctx.input_files([".pdf"]) or ctx.produced([".pdf"]))
    if len(files) < minimum:
        raise PluginError(f"Upload {'a PDF' if minimum == 1 else f'at least {minimum} PDFs'} for this tool.")
    return files


def _images(ctx):
    files = ctx.input_files(tk.IMAGE_EXT) or ctx.produced([".png", ".jpg", ".jpeg", ".webp"])
    if not files:
        raise PluginError("Upload one or more images for this tool.")
    return files


MAX_RENDER_PIXELS = 36_000_000


def render_page(page, dpi: int):
    """Render a PDF page at `dpi`, scaled down if the result would exceed MAX_RENDER_PIXELS."""
    w, h = page.get_size()  # points
    scale = dpi / 72
    if w * h * scale * scale > MAX_RENDER_PIXELS:
        scale = (MAX_RENDER_PIXELS / max(1.0, w * h)) ** 0.5
    return page.render(scale=scale).to_pil().convert("RGB")


def _chain(ctx, paths):
    ctx.state["pdf_chain"] = [{"name": Path(p).name, "path": str(p)} for p in paths]


def parse_pages(spec: str, total: int) -> list:
    pages = []
    for part in re.split(r"[,\s]+", spec or ""):
        if not part:
            continue
        m = re.fullmatch(r"(\d+)\s*-\s*(\d+|end|last)", part, re.I)
        if m:
            a = max(1, int(m.group(1)))
            b = total if m.group(2).lower() in ("end", "last") else min(total, int(m.group(2)))
            if a <= b:
                pages += list(range(a, b + 1))
        elif part.isdigit():
            pages.append(int(part))
    return [p for p in dict.fromkeys(pages) if 1 <= p <= total]


def _page_spec(problem: str) -> str:
    m = re.search(r"pages?\s+((?:\d+\s*(?:-\s*(?:\d+|end|last))?\s*,?\s*(?:and\s*)?)+)", problem, re.I)
    return m.group(1).replace("and", ",") if m else ""


def _zip(ctx, paths, name):
    zpath = ctx.out_path(name)
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        for p in paths:
            z.write(p, Path(p).name)
    return zpath


# ====================================================================== PDF -> Word
def pdf_to_docx(pdf_path, docx_path, log=lambda *_: None, check=lambda: None, max_pages: int = 300):
    import pdfplumber
    from docx import Document
    from docx.shared import Inches, Pt

    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    with pdfplumber.open(str(pdf_path)) as pdf:
        sizes = []
        for page in pdf.pages[:10]:
            sizes += [round(w["size"], 1) for w in page.extract_words(extra_attrs=["size"])]
        body = statistics.median(sizes) if sizes else 11
        for pno, page in enumerate(pdf.pages[:max_pages]):
            check()
            if pno:
                d.add_page_break()
            elements = []
            tables = page.find_tables()
            boxes = [t.bbox for t in tables]
            for t in tables:
                data = t.extract()
                if data and any(any(c for c in row) for row in data):
                    elements.append((t.bbox[1], "table", data))
            for im in page.images[:20]:
                try:
                    x0, top, x1, bottom = (max(0, im["x0"]), max(0, im["top"]), min(page.width, im["x1"]),
                                           min(page.height, im["bottom"]))
                    if (x1 - x0) < 40 or (bottom - top) < 40:
                        continue
                    pil = page.crop((x0, top, x1, bottom)).to_image(resolution=150).original
                    buf = io.BytesIO()
                    pil.convert("RGB").save(buf, "PNG")
                    elements.append((top, "image", (buf, (x1 - x0) / 72)))
                except Exception:
                    continue
            words = page.extract_words(extra_attrs=["size", "fontname"], use_text_flow=True, keep_blank_chars=False)
            words = [w for w in words if not any(b[0] - 1 <= w["x0"] <= b[2] + 1 and b[1] - 1 <= w["top"] <= b[3] + 1
                                                 for b in boxes)]
            lines = []
            for w in sorted(words, key=lambda w: (round(w["top"] / 3), w["x0"])):
                if lines and abs(lines[-1]["top"] - w["top"]) <= 3:
                    lines[-1]["words"].append(w)
                else:
                    lines.append({"top": w["top"], "words": [w]})
            para = None
            for ln in lines:
                ws = sorted(ln["words"], key=lambda w: w["x0"])
                text = " ".join(w["text"] for w in ws)
                size = statistics.mean(w["size"] for w in ws)
                bold = sum("bold" in w["fontname"].lower() for w in ws) > len(ws) / 2
                bottom = max(w["bottom"] for w in ws)
                is_head = size >= body * 1.25 and len(text) < 120
                bullet = re.match(r"^[•●▪◦\-–\*]\s+", text)
                new_para = (para is None or is_head or para["head"] or bullet or
                            ln["top"] - para["bottom"] > size * 0.9 or abs(size - para["size"]) > 1.5)
                if new_para:
                    if para:
                        elements.append((para["top"], "para", para))
                    para = {"top": ln["top"], "bottom": bottom, "text": text, "size": size, "bold": bold,
                            "head": is_head, "bullet": bool(bullet)}
                else:
                    para["text"] = (para["text"][:-1] if para["text"].endswith("-") else para["text"] + " ") + text
                    para["bottom"] = bottom
            if para:
                elements.append((para["top"], "para", para))
            for _, kind, payload in sorted(elements, key=lambda e: e[0]):
                if kind == "table":
                    ncol = max(len(r) for r in payload)
                    tbl = d.add_table(rows=len(payload), cols=ncol)
                    tbl.style = "Table Grid"
                    for ri, row in enumerate(payload):
                        for ci in range(ncol):
                            tbl.cell(ri, ci).text = (row[ci] if ci < len(row) and row[ci] else "").strip()
                    d.add_paragraph()
                elif kind == "image":
                    buf, width_in = payload
                    buf.seek(0)
                    d.add_picture(buf, width=Inches(min(6.3, max(1.0, width_in))))
                else:
                    if payload["head"]:
                        level = 1 if payload["size"] >= body * 1.6 else 2
                        d.add_heading(payload["text"], level=level)
                    elif payload["bullet"]:
                        d.add_paragraph(re.sub(r"^[•●▪◦\-–\*]\s+", "", payload["text"]), style="List Bullet")
                    else:
                        r = d.add_paragraph().add_run(payload["text"])
                        r.bold = payload["bold"] or None
        n_pages = len(pdf.pages)
    d.save(str(docx_path))
    return n_pages


@plugin("pdf_to_word", "PDF to Word", "pdf", "transform",
        "Converts PDFs into editable Word documents, keeping headings, paragraphs, bullet lists, tables and images.",
        icon="file-word", keywords=["pdf to word", "pdf to docx", "pdf into word", "editable", "convert pdf"],
        needs=["files:pdf"], ai_enhanced=False, example="Convert the attached PDF to Word")
def pdf_to_word(ctx):
    for f in _pdfs(ctx):
        out = ctx.out_path(Path(f["name"]).stem + ".docx")
        text_probe = tk.file_to_text(f["path"], max_chars=2000, ocr=False)
        if len(text_probe.strip()) < 40:
            ctx.log(f"{f['name']} looks scanned — running OCR first.", "warn")
            text = tk.file_to_text(f["path"], ocr=True, lang=_ocr_lang(ctx))
            from plg_docs import render_docx
            render_docx(tk.markdown_to_doc(text, Path(f["name"]).stem), out, {**_meta(ctx), "signature": False})
        else:
            n = pdf_to_docx(f["path"], out, check=ctx.check_cancel)
            ctx.log(f"Converted {f['name']} ({n} pages).")
        ctx.add_output(out, f"Word version of {f['name']}")


# ====================================================================== Office -> PDF
def _soffice():
    for name in ("soffice", "libreoffice"):
        p = shutil.which(name)
        if p:
            return p
    for p in (r"C:\Program Files\LibreOffice\program\soffice.exe", "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if os.path.exists(p):
            return p
    return None


@lru_cache(maxsize=1)
def _netns_prefix() -> list:
    """`unshare -rn` runs LibreOffice with no network at all (Linux, when the kernel allows it)."""
    if os.name == "nt" or not shutil.which("unshare"):
        return []
    try:
        ok = subprocess.run(["unshare", "-rn", "true"], capture_output=True, timeout=10).returncode == 0
    except (subprocess.SubprocessError, OSError):
        ok = False
    return ["unshare", "-rn"] if ok else []


OOXML = {".docx", ".pptx", ".xlsx", ".xlsm"}
ODF = {".odt", ".odp", ".ods"}
_FIELD_WORDS = re.compile(rb"(?i)\b(INCLUDEPICTURE|INCLUDETEXT|LINK|DDEAUTO|DDE|IMPORT)\b")
_REMOTE = rb"(?:https?|ftp|file|smb|webdav)://"


def _neutralise(name: str, data: bytes) -> bytes:
    """Remove only the things a converter would fetch; visible text and hyperlinks are untouched."""
    if name.endswith(".rels"):
        def rel(m):
            tag = m.group(0)
            if b'TargetMode="External"' in tag and b'/hyperlink"' not in tag:
                tag = re.sub(rb'Target="[^"]*"', b'Target="#"', tag)
            return tag
        return re.sub(rb"<Relationship\b[^>]*>", rel, data)
    fix = lambda m: m.group(1) + _FIELD_WORDS.sub(b"COMMENTS", m.group(2)) + m.group(3)
    data = re.sub(rb"(<w:instrText[^>]*>)([^<]*)(</w:instrText>)", fix, data)
    data = re.sub(rb'(w:instr=")([^"]*)(")', fix, data)
    data = re.sub(rb'(<(?:draw:image|draw:object|draw:object-ole|draw:plugin|draw:applet|text:section-source|'
                  rb'table:table-source|draw:floating-frame)\b[^>]*?xlink:href=")' + _REMOTE + rb'[^"]*"', rb'\1#"', data)
    if name.endswith("connections.xml"):
        data = re.sub(rb'\b(url|sourceFile|connection)="[^"]*"', rb'\1=""', data)
    return data


def sanitize_office(src: Path) -> bool:
    """Neutralise external references (linked images/media, remote templates, INCLUDEPICTURE fields, web
    queries) in an Office/ODF file. Returns False if the file can't be made safe."""
    if src.suffix.lower() not in OOXML | ODF:
        return False
    tmp = src.with_suffix(src.suffix + ".clean")
    try:
        with zipfile.ZipFile(src) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            total = 0
            for item in zin.infolist():
                total += item.file_size
                if total > 500 * 1024 * 1024:
                    raise zipfile.BadZipFile("too large when unpacked")
                data = zin.read(item.filename)
                if item.filename.endswith((".rels", ".xml")):
                    data = _neutralise(item.filename, data)
                zout.writestr(item, data)
        tmp.replace(src)
        return True
    except (zipfile.BadZipFile, OSError, KeyError, RuntimeError):
        tmp.unlink(missing_ok=True)
        return False


def office_convert(src, out_dir, target="pdf", timeout=180):
    exe = _soffice()
    if not exe:
        return None
    src = Path(src)
    isolated = _netns_prefix()
    # Without network isolation only sanitised OOXML/ODF files may reach LibreOffice.
    if not isolated and not sanitize_office(src):
        return None
    if isolated and src.suffix.lower() in OOXML | ODF:
        sanitize_office(src)
    env = dict(os.environ, http_proxy="http://127.0.0.1:9", https_proxy="http://127.0.0.1:9",
               HTTP_PROXY="http://127.0.0.1:9", HTTPS_PROXY="http://127.0.0.1:9", no_proxy="", NO_PROXY="")
    with tempfile.TemporaryDirectory(dir=str(config.TMP_DIR)) as profile:
        prof = Path(profile).as_posix() if os.name != "nt" else "/" + Path(profile).as_posix()
        cmd = isolated + [exe, f"-env:UserInstallation=file://{prof}", "--headless", "--norestore", "--nolockcheck",
                          "--convert-to", target, "--outdir", str(out_dir), str(src)]
        try:
            subprocess.run(cmd, capture_output=True, timeout=timeout, env=env)
        except (subprocess.SubprocessError, OSError):
            return None
    out = Path(out_dir) / (src.stem + "." + target.split(":")[0])
    return out if out.exists() and out.stat().st_size > 0 else None


@plugin("office_to_pdf", "Office to PDF", "pdf", "transform",
        "Converts Word, PowerPoint, Excel, text, Markdown and HTML files to PDF.",
        icon="file-pdf", keywords=["word to pdf", "docx to pdf", "ppt to pdf", "pptx to pdf", "excel to pdf",
                                   "xlsx to pdf", "convert to pdf", "doc to pdf", "txt to pdf"],
        needs=["files:office"], ai_enhanced=False, example="Convert my Word file to PDF")
def office_to_pdf(ctx):
    files = ctx.input_files(OFFICE_EXT)
    if not files:
        raise PluginError("Upload a Word, PowerPoint, Excel, text or HTML file to convert.")
    for f in files:
        ctx.check_cancel()
        out = None
        with tempfile.TemporaryDirectory(dir=str(config.TMP_DIR)) as tmp:
            src = Path(tmp) / f["name"]
            shutil.copyfile(f["path"], src)
            # plain text / markdown / html / csv never go through LibreOffice
            res = office_convert(src, tmp) if src.suffix.lower() not in (".txt", ".md", ".html", ".htm", ".csv") else None
            if res:
                out = ctx.out_path(Path(f["name"]).stem + ".pdf")
                shutil.copyfile(res, out)
        if out is None:
            if Path(f["name"]).suffix.lower() not in (".txt", ".md", ".html", ".htm", ".csv"):
                ctx.log(f"Using the built-in renderer for {f['name']} (layout simplified). Install LibreOffice "
                        "for pixel-faithful conversion.", "warn")
            text = tk.file_to_text(f["path"])
            doc = tk.markdown_to_doc(text, Path(f["name"]).stem)
            out = ctx.out_path(Path(f["name"]).stem + ".pdf")
            render_pdf(doc, out, _meta(ctx), plain=True)
        ctx.add_output(out, f"PDF of {f['name']}")
        ctx.state.setdefault("_office_pdfs", []).append(str(out))
        ctx.log(f"Converted {f['name']}.")
    _chain(ctx, ctx.state.pop("_office_pdfs", []))


# ====================================================================== merge / split / rotate
@plugin("merge_pdf", "Merge PDF", "pdf", "transform",
        "Combines multiple PDFs (and images) into one PDF in the order you uploaded them.",
        icon="merge", keywords=["merge", "combine", "join pdf", "merge pdf", "combine pdf"],
        needs=["files:pdf"], ai_enhanced=False, example="Merge these PDFs into one")
def merge_pdf(ctx):
    from pypdf import PdfReader, PdfWriter
    files = [f for f in ctx.inputs if Path(f["name"]).suffix.lower() in ({".pdf"} | tk.IMAGE_EXT)]
    if len(files) < 2:
        raise PluginError("Upload at least two PDFs (or images) to merge.")
    writer = PdfWriter()
    for f in files:
        if f["name"].lower().endswith(".pdf"):
            reader = PdfReader(f["path"])
            if reader.is_encrypted:
                raise PluginError(f"{f['name']} is password-protected. Unlock it first.")
            writer.append(reader)
        else:
            buf = io.BytesIO()
            with Image.open(f["path"]) as im:
                ImageOps.exif_transpose(im).convert("RGB").save(buf, "PDF", resolution=150)
            buf.seek(0)
            writer.append(PdfReader(buf))
    out = ctx.out_path("merged.pdf")
    with open(out, "wb") as fh:
        writer.write(fh)
    ctx.add_output(out, f"Merged PDF · {len(writer.pages)} pages")
    _chain(ctx, [out])
    ctx.log(f"Merged {len(files)} files into {len(writer.pages)} pages.")


@plugin("split_pdf", "Split PDF", "pdf", "transform",
        "Splits a PDF into single pages or the page ranges you choose (e.g. “pages 1-3, 5”).",
        icon="split", keywords=["split", "extract pages", "separate pages", "split pdf", "pages"],
        needs=["files:pdf"], params={"pages": "string"}, ai_enhanced=False, example="Extract pages 2-5 from this PDF")
def split_pdf(ctx, pages: str = ""):
    from pypdf import PdfReader, PdfWriter
    for f in _pdfs(ctx):
        reader = PdfReader(f["path"])
        if reader.is_encrypted:
            raise PluginError(f"{f['name']} is password-protected. Unlock it first.")
        total = len(reader.pages)
        spec = pages or _page_spec(ctx.problem)
        stem = Path(f["name"]).stem
        outs = []
        if spec:
            chosen = parse_pages(spec, total)
            if not chosen:
                raise PluginError(f"No valid pages in “{spec}” — the PDF has {total} pages.")
            w = PdfWriter()
            for p in chosen:
                w.add_page(reader.pages[p - 1])
            out = ctx.out_path(f"{stem}-pages-{tk.slugify(spec, 20)}.pdf")
            with open(out, "wb") as fh:
                w.write(fh)
            ctx.add_output(out, f"Pages {spec} of {f['name']}")
            continue
        for i, page in enumerate(reader.pages, 1):
            w = PdfWriter()
            w.add_page(page)
            out = ctx.job_dir / f"_{stem}-page-{i:03d}.pdf"
            with open(out, "wb") as fh:
                w.write(fh)
            outs.append(out)
        z = _zip(ctx, outs, f"{stem}-pages.zip")
        for o in outs:
            o.unlink(missing_ok=True)
        ctx.add_output(z, f"{total} single-page PDFs (ZIP)")
        ctx.log(f"Split {f['name']} into {total} pages.")


@plugin("rotate_pdf", "Rotate PDF", "pdf", "transform",
        "Rotates all pages or selected pages by 90°, 180° or 270°.",
        icon="rotate", keywords=["rotate", "turn pages", "landscape", "upside down"],
        needs=["files:pdf"], params={"degrees": "90|180|270", "pages": "string"}, ai_enhanced=False,
        example="Rotate pages 1-2 of this PDF by 90 degrees")
def rotate_pdf(ctx, degrees: int = 0, pages: str = ""):
    from pypdf import PdfReader, PdfWriter
    low = ctx.problem.lower()
    if not degrees:
        m = re.search(r"\b(90|180|270)\b", low)
        degrees = int(m.group(1)) if m else (270 if "left" in low or "anticlockwise" in low or "counter" in low else
                                             180 if "upside" in low else 90)
    outs = []
    for f in _pdfs(ctx):
        reader = PdfReader(f["path"])
        writer = PdfWriter(clone_from=reader)
        chosen = parse_pages(pages or _page_spec(ctx.problem), len(writer.pages)) or range(1, len(writer.pages) + 1)
        for p in chosen:
            writer.pages[p - 1].rotate(int(degrees))
        out = ctx.out_path(Path(f["name"]).stem + "-rotated.pdf")
        with open(out, "wb") as fh:
            writer.write(fh)
        ctx.add_output(out, f"Rotated {degrees}°")
        outs.append(out)
    _chain(ctx, outs)


# ====================================================================== compress
@plugin("compress_pdf", "Compress PDF", "pdf", "transform",
        "Reduces PDF size by recompressing images and streams (levels: low, medium, high).",
        icon="compress", keywords=["compress", "reduce size", "smaller", "shrink", "optimize pdf", "compress pdf"],
        needs=["files:pdf"], params={"level": "low|medium|high"}, ai_enhanced=False,
        example="Compress this PDF to make it smaller for e-mail")
def compress_pdf(ctx, level: str = ""):
    from pypdf import PdfReader, PdfWriter
    low = ctx.problem.lower()
    level = level or ("high" if any(w in low for w in ("maximum", "extreme", "high", "a lot", "very small")) else
                      "low" if "light" in low or "low" in low else "medium")
    quality, max_dim = {"low": (82, 2400), "medium": (62, 1700), "high": (42, 1200)}[level]
    outs = []
    for f in _pdfs(ctx):
        before = os.path.getsize(f["path"])
        writer = PdfWriter(clone_from=PdfReader(f["path"]))
        for page in writer.pages:
            for img in page.images:
                try:
                    pil = img.image
                    if pil.mode in ("RGBA", "LA", "P"):
                        continue  # keep transparency intact
                    if max(pil.size) > max_dim:
                        pil.thumbnail((max_dim, max_dim))
                    img.replace(pil.convert("RGB") if pil.mode not in ("RGB", "L") else pil, quality=quality)
                except Exception:
                    continue
            page.compress_content_streams()
        try:
            writer.compress_identical_objects(remove_identicals=True, remove_orphans=True)
        except Exception:
            pass
        out = ctx.out_path(Path(f["name"]).stem + "-compressed.pdf")
        with open(out, "wb") as fh:
            writer.write(fh)
        after = os.path.getsize(out)
        if after >= before:
            shutil.copyfile(f["path"], out)
            ctx.log(f"{f['name']} is already well optimised — no further reduction possible without quality loss.",
                    "warn")
            after = before
        ctx.add_output(out, f"Compressed · {before / 1e6:.2f} MB → {after / 1e6:.2f} MB")
        outs.append(out)
        ctx.log(f"{f['name']}: {before / 1e6:.2f} MB → {after / 1e6:.2f} MB ({(1 - after / before) * 100:.0f}% smaller).")
    _chain(ctx, outs)


# ====================================================================== watermark / numbers
def _overlay(page_w, page_h, draw_fn):
    from pypdf import PdfReader
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(page_w, page_h))
    draw_fn(c, page_w, page_h)
    c.save()
    buf.seek(0)
    return PdfReader(buf).pages[0]


@plugin("watermark_pdf", "Watermark PDF", "pdf", "transform",
        "Stamps a diagonal text watermark (e.g. CONFIDENTIAL, DRAFT or your brand) on every page.",
        icon="stamp", keywords=["watermark", "stamp", "confidential", "draft"], needs=["files:pdf"],
        params={"text": "string"}, ai_enhanced=False, example='Add watermark "DRAFT" to this PDF')
def watermark_pdf(ctx, text: str = ""):
    from pypdf import PdfReader, PdfWriter
    if not text:
        m = re.search(r"[\"“']([^\"”']{1,40})[\"”']", ctx.problem) or re.search(
            r"watermark(?:\s+(?:text|saying|with|as|of))?\s+([A-Za-z0-9 &.\-]{2,40}?)(?:\s+(?:to|on|in)\b|$)", ctx.problem, re.I)
        text = m.group(1).strip() if m else "CONFIDENTIAL"
    fnt = brand.reportlab_fonts()["bold"]
    outs = []
    for f in _pdfs(ctx):
        writer = PdfWriter(clone_from=PdfReader(f["path"]))
        cache = {}
        for page in writer.pages:
            w, h = float(page.mediabox.width), float(page.mediabox.height)
            if (w, h) not in cache:
                def draw(c, w, h):
                    c.saveState()
                    c.setFillColorRGB(0.55, 0.45, 0.25)
                    c.setFillAlpha(0.16)
                    size = min(w, h) / max(4, len(text) * 0.62)
                    c.setFont(fnt, size)
                    c.translate(w / 2, h / 2)
                    c.rotate(35)
                    c.drawCentredString(0, -size / 3, brand.pdf_safe(text))
                    c.restoreState()
                cache[(w, h)] = _overlay(w, h, draw)
            page.merge_page(cache[(w, h)])
        out = ctx.out_path(Path(f["name"]).stem + "-watermarked.pdf")
        with open(out, "wb") as fh:
            writer.write(fh)
        ctx.add_output(out, f"Watermarked “{text}”")
        outs.append(out)
    _chain(ctx, outs)


@plugin("page_numbers", "Add Page Numbers", "pdf", "transform",
        "Adds “Page X of Y” numbering to the bottom of every page.",
        icon="hash", keywords=["page number", "page numbers", "number the pages", "pagination"], needs=["files:pdf"],
        ai_enhanced=False, example="Add page numbers to this PDF")
def page_numbers(ctx):
    from pypdf import PdfReader, PdfWriter
    fnt = brand.reportlab_fonts()["regular"]
    outs = []
    for f in _pdfs(ctx):
        writer = PdfWriter(clone_from=PdfReader(f["path"]))
        total = len(writer.pages)
        for i, page in enumerate(writer.pages, 1):
            w, h = float(page.mediabox.width), float(page.mediabox.height)

            def draw(c, w, h, i=i):
                c.setFont(fnt, 9)
                c.setFillColorRGB(0.35, 0.33, 0.3)
                c.drawCentredString(w / 2, 18, f"Page {i} of {total}")
            page.merge_page(_overlay(w, h, draw))
        out = ctx.out_path(Path(f["name"]).stem + "-numbered.pdf")
        with open(out, "wb") as fh:
            writer.write(fh)
        ctx.add_output(out, "PDF with page numbers")
        outs.append(out)
    _chain(ctx, outs)


# ====================================================================== protect / unlock
@plugin("protect_pdf", "Protect PDF", "pdf", "transform",
        "Encrypts a PDF with a password (AES-256). Say: “protect with password XYZ”.",
        icon="lock", keywords=["protect", "encrypt", "password protect", "lock pdf", "secure pdf"], needs=["files:pdf"],
        params={"password": "secret"}, ai_enhanced=False, example="Protect this PDF with password Ruby@2026")
def protect_pdf(ctx, password: str = ""):
    from pypdf import PdfReader, PdfWriter
    if not password or len(password) < 4:
        raise PluginError("Give a password of at least 4 characters, e.g. “protect with password Ruby@2026”.")
    for f in _pdfs(ctx):
        writer = PdfWriter(clone_from=PdfReader(f["path"]))
        writer.encrypt(user_password=password, owner_password=None, algorithm="AES-256")
        out = ctx.out_path(Path(f["name"]).stem + "-protected.pdf")
        with open(out, "wb") as fh:
            writer.write(fh)
        ctx.add_output(out, "Password-protected PDF (AES-256)")
    ctx.log("Encrypted. The password is not stored anywhere — keep it safe.")


@plugin("unlock_pdf", "Unlock PDF", "pdf", "transform",
        "Removes the password from a PDF you own, using the password you provide (no cracking).",
        icon="unlock", keywords=["unlock", "remove password", "decrypt"], needs=["files:pdf"],
        params={"password": "secret"}, ai_enhanced=False, example="Unlock this PDF, the password is 1234")
def unlock_pdf(ctx, password: str = ""):
    from pypdf import PdfReader, PdfWriter
    outs = []
    for f in _pdfs(ctx):
        reader = PdfReader(f["path"])
        if not reader.is_encrypted:
            ctx.log(f"{f['name']} is not encrypted.", "warn")
            continue
        if not password or not reader.decrypt(password):
            raise PluginError("Wrong or missing password. Write: “unlock this PDF, password is …”.")
        writer = PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        out = ctx.out_path(Path(f["name"]).stem + "-unlocked.pdf")
        with open(out, "wb") as fh:
            writer.write(fh)
        ctx.add_output(out, "Unlocked PDF")
        outs.append(out)
    _chain(ctx, outs)


# ====================================================================== images <-> PDF
@plugin("pdf_to_images", "PDF to Images", "pdf", "transform",
        "Renders PDF pages to high-quality PNG or JPG images.",
        icon="image", keywords=["pdf to jpg", "pdf to png", "pdf to image", "pdf to images", "convert pdf to image"],
        needs=["files:pdf"], params={"format": "png|jpg", "dpi": "int"}, ai_enhanced=False,
        example="Convert this PDF to JPG images")
def pdf_to_images(ctx, format: str = "", dpi: int = 0):
    import pypdfium2 as pdfium
    low = ctx.problem.lower()
    fmt = format or ("jpg" if ("jpg" in low or "jpeg" in low) else "png")
    dpi = int(dpi or (re.search(r"(\d{2,3})\s*dpi", low).group(1) if re.search(r"(\d{2,3})\s*dpi", low) else 150))
    dpi = max(50, min(400, dpi))
    for f in _pdfs(ctx):
        pdf = pdfium.PdfDocument(f["path"])
        paths = []
        try:
            for i in range(min(len(pdf), 300)):
                ctx.check_cancel()
                img = render_page(pdf[i], dpi)
                p = ctx.job_dir / f"{Path(f['name']).stem}-page-{i + 1:03d}.{fmt}"
                img.save(p, "JPEG" if fmt == "jpg" else "PNG", quality=90)
                paths.append(p)
        finally:
            pdf.close()
        for p in paths[:6]:
            ctx.add_output(p, f"Page image {p.stem.rsplit('-', 1)[-1]}")
        if len(paths) > 1:
            ctx.add_output(_zip(ctx, paths, f"{Path(f['name']).stem}-images.zip"), f"All {len(paths)} pages (ZIP)")
        ctx.log(f"Rendered {len(paths)} page(s) at {dpi} DPI.")


@plugin("images_to_pdf", "Images to PDF", "pdf", "transform",
        "Combines photos or scans into one A4 PDF, auto-rotated and fitted to the page.",
        icon="file-pdf", keywords=["jpg to pdf", "png to pdf", "image to pdf", "images to pdf", "photos to pdf", "scan to pdf"],
        needs=["files:image"], ai_enhanced=False, example="Make a PDF from these photos")
def images_to_pdf(ctx):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    files = _images(ctx)
    out = ctx.out_path("images.pdf")
    c = canvas.Canvas(str(out), pagesize=A4)
    for f in files:
        with Image.open(f["path"]) as im:
            im = ImageOps.exif_transpose(im).convert("RGB")
            size = landscape(A4) if im.width > im.height else A4
            c.setPageSize(size)
            margin = 24
            avail_w, avail_h = size[0] - 2 * margin, size[1] - 2 * margin
            scale = min(avail_w / im.width, avail_h / im.height)
            w, h = im.width * scale, im.height * scale
            c.drawImage(ImageReader(im), (size[0] - w) / 2, (size[1] - h) / 2, w, h)
            c.showPage()
    c.save()
    ctx.add_output(out, f"PDF from {len(files)} image(s)")
    _chain(ctx, [out])


# ====================================================================== OCR / text
def _ocr_lang(ctx) -> str:
    low = ctx.problem.lower()
    wanted = ["eng"]
    for word, code in (("hindi", "hin"), ("kannada", "kan"), ("tamil", "tam"), ("telugu", "tel"), ("marathi", "mar"),
                       ("malayalam", "mal"), ("bengali", "ben"), ("gujarati", "guj"), ("french", "fra"),
                       ("german", "deu"), ("spanish", "spa"), ("arabic", "ara")):
        if word in low:
            wanted.append(code)
    try:
        import pytesseract
        have = set(pytesseract.get_languages(config=""))
        wanted = [w for w in wanted if w in have] or ["eng"]
    except Exception:
        pass
    return "+".join(wanted)


@plugin("ocr_extract", "OCR — Image/Scan to Text", "pdf", "transform",
        "Reads text from photos, screenshots and scanned PDFs, and creates a searchable PDF.",
        icon="scan", keywords=["ocr", "scan", "scanned", "extract text from image", "read text", "image to text",
                               "searchable pdf", "handwriting"],
        needs=["files:image|pdf"], ai_enhanced=False, example="Extract the text from this scanned PDF")
def ocr_extract(ctx):
    import pytesseract
    from pypdf import PdfReader, PdfWriter
    try:
        pytesseract.get_tesseract_version()
    except Exception:
        raise PluginError("OCR engine (Tesseract) is not installed on this server.")
    lang = _ocr_lang(ctx)
    ctx.state["ocr_lang"] = lang
    files = ctx.input_files(tk.IMAGE_EXT | {".pdf"})
    if not files:
        raise PluginError("Upload images or a scanned PDF.")
    texts, writer = [], PdfWriter()

    def pages_of(f):
        """Yield one page image at a time so memory stays flat."""
        if f["name"].lower().endswith(".pdf"):
            import pypdfium2 as pdfium
            pdf = pdfium.PdfDocument(f["path"])
            try:
                for i in range(min(len(pdf), 60)):
                    yield render_page(pdf[i], 300)
            finally:
                pdf.close()
        else:
            with Image.open(f["path"]) as im:
                im = ImageOps.exif_transpose(im).convert("RGB")
                if im.width * im.height > MAX_RENDER_PIXELS:
                    im.thumbnail((8000, 8000))
                yield im

    for f in files:
        for n, img in enumerate(pages_of(f), 1):
            ctx.check_cancel()
            ctx.log(f"Reading {f['name']} page {n}")
            texts.append(f"## {f['name']} — page {n}\n\n" + pytesseract.image_to_string(img, lang=lang).strip())
            pdf_bytes = pytesseract.image_to_pdf_or_hocr(img, extension="pdf", lang=lang)
            writer.append(PdfReader(io.BytesIO(pdf_bytes)))
            img.close()
    full = "\n\n".join(texts)
    ctx.state["input_text"] = full
    ctx.add_text_output("ocr-text.txt", full, "Extracted text", "text")
    out = ctx.out_path("searchable.pdf")
    with open(out, "wb") as fh:
        writer.write(fh)
    ctx.add_output(out, "Searchable PDF")
    words = len(full.split())
    ctx.log(f"OCR complete ({lang}): {words:,} words.")


@plugin("extract_text", "Extract Text", "pdf", "transform",
        "Pulls plain text out of PDF, Word, PowerPoint, Excel, HTML and other documents.",
        icon="text", keywords=["extract text", "get text", "text from", "copy text", "to txt", "plain text"],
        needs=["files:any"], ai_enhanced=False, example="Extract all text from this PowerPoint")
def extract_text(ctx):
    if not ctx.inputs:
        raise PluginError("Upload a document to extract text from.")
    for f in ctx.inputs:
        text = tk.file_to_text(f["path"], lang=_ocr_lang(ctx))
        ctx.add_text_output(Path(f["name"]).stem + ".txt", text, f"Text of {f['name']}", "text")


# ====================================================================== image tools
PRESETS = {"instagram post": (1080, 1350), "instagram square": (1080, 1080), "instagram story": (1080, 1920),
           "story": (1080, 1920), "reel": (1080, 1920), "youtube thumbnail": (1280, 720), "thumbnail": (1280, 720),
           "facebook cover": (1640, 624), "linkedin banner": (1584, 396), "whatsapp dp": (640, 640),
           "passport": (413, 531), "hd": (1920, 1080), "4k": (3840, 2160)}


@plugin("image_tools", "Image Converter & Resizer", "media", "transform",
        "Resizes, crops to social sizes, converts (JPG/PNG/WEBP), compresses and strips private metadata from images.",
        icon="image", keywords=["resize", "crop", "convert image", "compress image", "jpg", "png", "webp", "reduce image",
                                "thumbnail", "remove metadata", "exif"],
        needs=["files:image"], params={"width": "int", "height": "int", "format": "jpg|png|webp", "quality": "int"},
        ai_enhanced=False, example="Resize these photos for an Instagram story and convert to JPG")
def image_tools(ctx, width: int = 0, height: int = 0, format: str = "", quality: int = 0):
    low = ctx.problem.lower()
    files = _images(ctx)
    size = None
    for name, dims in PRESETS.items():
        if name in low:
            size = dims
            break
    m = re.search(r"(\d{2,5})\s*[x×*]\s*(\d{2,5})", low)
    if width and height:
        size = (int(width), int(height))
    elif m:
        size = (int(m.group(1)), int(m.group(2)))
    if size:
        size = (max(16, min(8192, size[0])), max(16, min(8192, size[1])))
    pct = re.search(r"(\d{1,3})\s*%", low)
    fmt = format or next((f for f in ("webp", "png", "jpg", "jpeg") if re.search(r"\b" + f + r"\b", low)), "")
    fmt = "jpg" if fmt == "jpeg" else fmt
    quality = int(quality or (55 if "compress" in low or "reduce" in low or "smaller" in low else 88))
    outs = []
    for f in files:
        with Image.open(f["path"]) as im:
            im = ImageOps.exif_transpose(im)
            if size:
                im = ImageOps.fit(im.convert("RGBA") if im.mode in ("RGBA", "LA", "P") else im.convert("RGB"),
                                  size, Image.LANCZOS, centering=(0.5, 0.45))
            elif pct and not ("quality" in low):
                k = max(1, min(400, int(pct.group(1)))) / 100
                k = min(k, 8192 / max(im.width, im.height))
                im = im.resize((max(1, int(im.width * k)), max(1, int(im.height * k))), Image.LANCZOS)
            target = fmt or (Path(f["name"]).suffix.lower().lstrip(".") or "png")
            target = "jpg" if target == "jpeg" else target
            if target not in ("jpg", "png", "webp"):
                target = "png"
            out = ctx.out_path(Path(f["name"]).stem + f"-edited.{target}")
            if target == "jpg":
                bg = Image.new("RGB", im.size, (255, 255, 255))
                rgba = im.convert("RGBA")
                bg.paste(rgba, mask=rgba.getchannel("A"))
                bg.save(out, "JPEG", quality=quality, optimize=True, progressive=True)
            elif target == "webp":
                im.save(out, "WEBP", quality=quality, method=5)
            else:
                im.save(out, "PNG", optimize=True)
            outs.append(out)
            ctx.add_output(out, f"{im.width}×{im.height} {target.upper()} (metadata removed)")
    if len(outs) > 3:
        ctx.add_output(_zip(ctx, outs, "images.zip"), f"All {len(outs)} images (ZIP)")
    ctx.log(f"Processed {len(outs)} image(s).")


@plugin("qr_generator", "QR Code Generator", "media", "render",
        "Creates high-resolution QR codes for links, UPI IDs, Wi-Fi, phone numbers or any text.",
        icon="qr", keywords=["qr", "qr code", "qrcode", "scan code"], params={"data": "string", "password": "secret"},
        ai_enhanced=False,
        example="Make a QR code for https://inera.example.com")
def qr_generator(ctx, data: str = "", password: str = ""):
    import qrcode
    from qrcode.constants import ERROR_CORRECT_H
    if not data:
        url = re.search(r"https?://\S+", ctx.problem)
        upi = re.search(r"\b[\w.\-]{2,}@[a-z]{2,}\b", ctx.problem)
        quoted = re.search(r"[\"“']([^\"”']{1,500})[\"”']", ctx.problem)
        wifi = re.search(r"wi-?fi\s+(?:name|ssid)?\s*[:=]?\s*(\S+)\s+(?:password|pass)\s*[:=]?\s*(\S+)", ctx.problem, re.I)
        if wifi:
            pw = password if (password and "•" in wifi.group(2)) else wifi.group(2)
            data = f"WIFI:T:WPA;S:{wifi.group(1)};P:{pw};;"
        elif url:
            data = url.group(0).rstrip(".,)")
        elif upi and "upi" in ctx.problem.lower():
            data = f"upi://pay?pa={upi.group(0)}&pn={ctx.business}"
        elif quoted:
            data = quoted.group(1)
        else:
            m = re.search(r"\bfor\s+(.+)$", ctx.problem, re.I)
            data = (m.group(1) if m else ctx.problem).strip()
    qr = qrcode.QRCode(error_correction=ERROR_CORRECT_H, box_size=16, border=3)
    qr.add_data(data)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white").convert("RGB")
    out = ctx.out_path("qr-code.png")
    img.save(out)
    ctx.add_output(out, "QR code")
    ctx.log(f"QR code encodes: {data[:80]}{'…' if len(data) > 80 else ''} — scan-test it before printing.")


@plugin("zip_files", "ZIP Files", "pdf", "transform",
        "Bundles your uploaded files (or this run's outputs) into a single ZIP archive.",
        icon="zip", keywords=["zip", "archive", "bundle files", "compress files"], ai_enhanced=False,
        example="Zip all these files together")
def zip_files(ctx):
    paths = [f["path"] for f in ctx.inputs] or [p["path"] for p in ctx.produced()]
    if not paths:
        raise PluginError("Nothing to zip — upload files first.")
    ctx.add_output(_zip(ctx, paths, "files.zip"), f"ZIP of {len(paths)} file(s)")
