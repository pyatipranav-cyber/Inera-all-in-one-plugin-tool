"""Offline text intelligence: sentence splitting, keywords, extractive
summaries, quiz/flashcard generation, the shared document model and
file-to-text extraction. No network, no AI engine required."""
import csv
import html
import io
import math
import random
import re
from collections import Counter
from html.parser import HTMLParser
from pathlib import Path

STOPWORDS = set("""
a about above after again against all also am an and any are aren't as at be because been before being
below between both but by can can't cannot could couldn't did didn't do does doesn't doing don't down during
each few for from further get got had hadn't has hasn't have haven't having he he'd he'll he's her here here's
hers herself him himself his how how's however i i'd i'll i'm i've if in into is isn't it it's its itself just
let's like made make many may me might more most much must mustn't my myself new no nor not now of off often on
once one only or other ought our ours ourselves out over own per same shan't she she'd she'll she's should
shouldn't since so some such than that that's the their theirs them themselves then there there's these they
they'd they'll they're they've this those though through thus to too two under until up upon us use used using
very via was wasn't we we'd we'll we're we've were weren't what what's when when's where where's whether which
while who who's whom why why's will with within without won't would wouldn't yet you you'd you'll you're
you've your yours yourself yourselves also among around called known include including includes several
well first second third later early another every either neither within became become becomes would could
""".split())

_ABBREV = {"mr", "mrs", "ms", "dr", "prof", "sr", "jr", "st", "vs", "etc", "e.g", "i.e", "inc", "ltd",
           "co", "no", "fig", "approx", "dept", "est", "u.s", "u.k", "jan", "feb", "mar", "apr", "jun",
           "jul", "aug", "sep", "sept", "oct", "nov", "dec"}


# ---------------------------------------------------------------- basics
def normalize_ws(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


HEADING_LINE = re.compile(r"^\s*(?:=+[^=\n]+=+|#{1,6}\s.*|# SOURCE:.*)\s*$", re.M)


def split_sentences(text: str) -> list:
    """Sentence splitter that respects paragraph/line boundaries and ignores heading lines."""
    text = HEADING_LINE.sub("\n", text or "")
    out = []
    for block in re.split(r"\n\s*\n|\n(?=\s*(?:[-*•]|\d+[.)])\s)", text):
        for line in _split_block(block):
            out.append(line)
    return out


def _split_block(text: str) -> list:
    text = re.sub(r"^\s*(?:[-*•]|\d+[.)])\s+", "", text or "")
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return []
    parts = re.split(r"(?<=[.!?])\s+(?=[\"'(\[]?[A-Z0-9])", text)
    out, buf = [], []
    for part in parts:
        buf.append(part)
        tail = part.rsplit(" ", 1)[-1]
        m = re.fullmatch(r"([A-Za-z.]+)\.", tail)
        if m and m.group(1).lower().rstrip(".") in _ABBREV and len(buf) < 50:
            continue  # abbreviation such as "Dr." — keep joining
        out.append(" ".join(buf))
        buf = []
    if buf:
        out.append(" ".join(buf))
    return [s.strip() for s in out if len(s.strip()) > 1]


def words(text: str) -> list:
    return re.findall(r"[A-Za-z][A-Za-z\-']+|\d+(?:\.\d+)?", text or "")


def content_words(text: str) -> list:
    return [w.lower() for w in words(text) if w.lower() not in STOPWORDS and len(w) > 2 and not w.isdigit()]


def truncate_words(text: str, n: int) -> str:
    ws = (text or "").split()
    return " ".join(ws[:n]) + ("…" if len(ws) > n else "")


def strip_md(text: str) -> str:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text or "")
    text = re.sub(r"(?<!\w)\*(?!\s)(.+?)(?<!\s)\*(?!\w)", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1 (\2)", text)
    return text


def md_runs(text: str) -> list:
    """Split text with **bold** / *italic* markers into (text, bold, italic) runs."""
    runs = []
    pattern = re.compile(r"(\*\*.+?\*\*|(?<!\w)\*(?!\s).+?(?<!\s)\*(?!\w))")
    pos = 0
    for m in pattern.finditer(text or ""):
        if m.start() > pos:
            runs.append((text[pos:m.start()], False, False))
        token = m.group(0)
        if token.startswith("**"):
            runs.append((token[2:-2], True, False))
        else:
            runs.append((token[1:-1], False, True))
        pos = m.end()
    if pos < len(text or ""):
        runs.append((text[pos:], False, False))
    return runs or [("", False, False)]


def slugify(text: str, maxlen: int = 50) -> str:
    s = re.sub(r"[^a-zA-Z0-9]+", "-", (text or "").lower()).strip("-")
    return (s[:maxlen].strip("-") or "output")


def title_case(text: str) -> str:
    small = {"a", "an", "the", "and", "or", "of", "in", "on", "for", "to", "with", "at", "by", "vs"}
    ws = (text or "").split()
    out = []
    for i, w in enumerate(ws):
        if w.isupper() and len(w) > 1:
            out.append(w)
        elif i and w.lower() in small:
            out.append(w.lower())
        else:
            out.append(w[:1].upper() + w[1:])
    return " ".join(out)


# ---------------------------------------------------------------- keywords
VERB_BREAKERS = set("""
show shows showed shown demonstrate demonstrates demonstrated propose proposes proposed estimate estimates estimated
produce produces produced use uses turn turns turned fix fixes fixed capture captures captured convert converts
converted help helps helped take takes took taken find finds found reduce reduces reduced increase increases
increased contain contains contained allow allows allowed begin begins began provide provides provided require
requires required describe describes described consider considered remain remains remained occur occurs occurred
work works worked carry carries carried form forms formed release releases released absorb absorbs absorbed
reflect reflects reflected collect collects collected save saves saved study studies studied lead leads led
grow grows grew grown concluded conclude concludes suggest suggests suggested receive receives received create
creates created build builds built keep keeps kept give gives gave given need needs needed want wants wanted
say says said tell tells told see sees saw seen get gets got come comes came go goes went gone run runs ran
look looks looked appear appears appeared continue continues continued explain explains explained tend tends
seem seems seemed start starts started stop stops stopped write writes wrote written learn learns learned
reach reaches reached open opens opened close closes closed regenerate regenerates attaches attach attached
traced trace traces earned earn earns pump pumps pumped travel travels travelled traveled split splits limit
limits limited slow slows slowed raise raises raised accumulate accumulates accumulated regulate regulates
regulated purify purifies purified restore restores restored supply supplies supplied breathe breathes breathed
studying using making taking helping producing including becoming having being getting giving showing
almost mainly mostly often usually finally largely nearly generally typically particularly especially really simply
directly widely highly sometimes always never already still even just quite rather fairly much
""".split())


def keywords(text: str, n: int = 15) -> list:
    """RAKE-style key phrase extraction (1-3 words) with frequency scoring."""
    sentences = split_sentences(text)
    phrase_counts = Counter()
    display = {}
    word_freq, word_deg = Counter(), Counter()
    for sent in sentences:
        tokens = re.findall(r"[A-Za-z][A-Za-z\-']*|\d+", sent)
        phrase = []
        for tok in tokens + ["."]:
            low = tok.lower()
            if low in STOPWORDS or low in VERB_BREAKERS or tok == "." or len(tok) < 3 or tok.isdigit() or \
                    (tok.islower() and len(tok) > 5 and low.endswith("ed")):
                if 1 <= len(phrase) <= 3:
                    key = " ".join(w.lower() for w in phrase)
                    phrase_counts[key] += 1
                    display.setdefault(key, " ".join(phrase))
                    for w in phrase:
                        word_freq[w.lower()] += 1
                        word_deg[w.lower()] += len(phrase)
                phrase = []
            else:
                phrase.append(tok)
    if not phrase_counts:
        return []
    scores = {}
    for key, cnt in phrase_counts.items():
        ws = key.split()
        score = sum((word_deg[w] + word_freq[w]) / max(word_freq[w], 1) for w in ws)
        score *= math.log(1 + cnt) + (0.6 if cnt > 1 else 0)
        if display[key][:1].isupper():
            score *= 1.15
        scores[key] = score
    ranked = sorted(scores, key=lambda k: -scores[k])
    out, seen_words = [], set()
    for key in ranked:
        if len(key) < 4:
            continue
        ws = set(key.split())
        if ws <= seen_words and len(ws) == 1:
            continue
        out.append(display[key])
        seen_words |= ws
        if len(out) >= n:
            break
    return out


# ---------------------------------------------------------------- summary
def rank_sentences(text: str) -> list:
    sentences = [s for s in split_sentences(text) if 5 <= len(s.split()) <= 60]
    if not sentences:
        return []
    freq = Counter(content_words(text))
    if not freq:
        return [(1.0, i, s) for i, s in enumerate(sentences)]
    top = freq.most_common(1)[0][1]
    ranked = []
    for i, s in enumerate(sentences):
        cw = content_words(s)
        if not cw:
            continue
        score = sum(freq[w] / top for w in cw) / (len(cw) ** 0.55)
        score *= 1.0 + 0.35 * max(0, 3 - i) / 3  # lead bias
        ranked.append((score, i, s))
    return sorted(ranked, key=lambda t: -t[0])


def summarize(text: str, n: int = 5) -> list:
    ranked = rank_sentences(text)
    chosen = sorted(ranked[:n], key=lambda t: t[1])
    out, seen = [], set()
    for _, _, s in chosen:
        sig = " ".join(content_words(s)[:6])
        if sig in seen:
            continue
        seen.add(sig)
        out.append(s)
    return out


def summary_paragraph(text: str, n: int = 4) -> str:
    return " ".join(summarize(text, n))


# ---------------------------------------------------------------- sections
SKIP_SECTIONS = {"see also", "references", "external links", "notes", "further reading",
                 "bibliography", "sources", "citations", "footnotes", "gallery", "works cited"}


def split_wiki_sections(text: str) -> list:
    """Split Wikipedia plain-text extracts (== Heading ==) into (heading, body, level)."""
    out = []
    current, level, buf = "Introduction", 1, []
    for line in (text or "").splitlines():
        m = re.match(r"^(={2,6})\s*(.+?)\s*\1\s*$", line.strip())
        if m:
            if "".join(buf).strip():
                out.append((current, normalize_ws("\n".join(buf)), level))
            current, level, buf = m.group(2), len(m.group(1)) - 1, []
        else:
            buf.append(line)
    if "".join(buf).strip():
        out.append((current, normalize_ws("\n".join(buf)), level))
    return [s for s in out if s[0].lower() not in SKIP_SECTIONS]


def split_plain_sections(text: str, max_sections: int = 8) -> list:
    """Heuristic sectioning of plain text: markdown headings, short title lines, or chunks."""
    lines = (text or "").splitlines()
    sections, current, buf = [], "Overview", []
    for line in lines:
        s = line.strip()
        md = re.match(r"^#{1,4}\s+(.+)$", s)
        is_title = (md is not None) or (
            0 < len(s) < 70 and len(s.split()) <= 9 and not s.endswith((".", ",", ";", ":"))
            and s[:1].isupper() and sum(c.isalpha() for c in s) > 3 and buf and not buf[-1].strip())
        if is_title:
            if " ".join(buf).strip():
                sections.append((current, normalize_ws("\n".join(buf)), 1))
            current, buf = (md.group(1) if md else s), []
        else:
            buf.append(line)
    if " ".join(buf).strip():
        sections.append((current, normalize_ws("\n".join(buf)), 1))
    if len(sections) <= 1:
        sents = split_sentences(text)
        if len(sents) > 12:
            size = max(6, len(sents) // max_sections)
            sections = []
            for i in range(0, len(sents), size):
                chunk = " ".join(sents[i:i + size])
                kw = keywords(chunk, 1)
                sections.append((title_case(kw[0]) if kw else f"Part {len(sections) + 1}", chunk, 1))
    return sections[:max_sections * 2]


# ---------------------------------------------------------------- quiz / cards
def make_mcqs(text: str, n: int = 10, seed: int = 7) -> list:
    rng = random.Random(seed)
    kws = keywords(text, 60)
    ranked = rank_sentences(text)
    questions, used_answers, used_sents = [], set(), set()
    for _, _, sent in ranked:
        if len(questions) >= n:
            break
        if len(sent.split()) < 7 or sent in used_sents:
            continue
        year = re.search(r"\b(1[5-9]\d\d|20\d\d)\b", sent)
        answer = None
        if year and rng.random() < 0.5:
            answer = year.group(1)
            y = int(answer)
            deltas = rng.sample([-15, -9, -5, -3, 3, 4, 7, 11, 20], 3)
            distractors = [str(y + d) for d in deltas]
        else:
            for kw in sorted(kws, key=len, reverse=True):
                if kw.lower() in used_answers:
                    continue
                if re.search(r"\b" + re.escape(kw) + r"\b", sent) and not sent.lower().startswith(kw.lower()):
                    answer = kw
                    break
            if not answer:
                continue
            pool = [k for k in kws if k.lower() != answer.lower() and k.lower() not in sent.lower()
                    and abs(len(k.split()) - len(answer.split())) <= 1]
            if len(pool) < 3:
                continue
            same_case = [k for k in pool if k[:1].isupper() == answer[:1].isupper()]
            src = same_case if len(same_case) >= 3 else pool
            distractors = rng.sample(src[:15], 3) if len(src) >= 3 else src[:3]
        stem = re.sub(r"\b" + re.escape(answer) + r"\b", "_____", sent, count=1)
        options = distractors + [answer]
        rng.shuffle(options)
        questions.append({
            "question": "Fill in the blank: " + stem,
            "options": options,
            "answer": options.index(answer),
            "explanation": sent,
        })
        used_answers.add(answer.lower())
        used_sents.add(sent)
    return questions


_DEF_RE = re.compile(r"^(?:The |An |A )?(?P<term>[A-Z][\w\-']*(?: [\w\-']+){0,4}?)\s+(?:is|are|was|were)\s+"
                     r"(?P<def>(?:a|an|the|one|any|used|known|called|defined)\b.{12,300})$|"
                     r"^(?:The |An |A )?(?P<term2>[A-Z][\w\-']*(?: [\w\-']+){0,4}?)\s+(?:refers to|means|denotes)\s+"
                     r"(?P<def2>.{12,300})$")
_BAD_TERM_START = {"when", "if", "some", "most", "many", "this", "these", "those", "it", "there", "together", "they",
                   "he", "she", "we", "you", "such", "each", "both", "all", "one", "another", "while", "during",
                   "after", "before", "in", "on", "at", "for", "as", "however", "because", "since", "the", "its"}


def make_flashcards(text: str, n: int = 20) -> list:
    cards, seen = [], set()
    for sent in split_sentences(text):
        m = _DEF_RE.match(sent.strip())
        if m:
            term = (m.group("term") or m.group("term2") or "").strip()
            if term and term.split()[0].lower() not in _BAD_TERM_START and len(term.split()) <= 5 and \
                    term.lower() not in seen:
                seen.add(term.lower())
                cards.append({"front": term, "back": sent.strip()})
        if len(cards) >= n:
            return cards
    sentences = split_sentences(text)
    for kw in keywords(text, n * 2):
        if len(cards) >= n:
            break
        if kw.lower() in seen:
            continue
        for s in sentences:
            if kw.lower() in s.lower() and 6 <= len(s.split()) <= 45:
                seen.add(kw.lower())
                cards.append({"front": kw, "back": s})
                break
    return cards


# ---------------------------------------------------------------- doc model
def new_doc(title: str, subtitle: str = "") -> dict:
    return {"title": title, "subtitle": subtitle, "sections": [], "sources": []}


def add_section(doc: dict, heading: str, paragraphs=None, bullets=None, numbered=None, table=None) -> dict:
    blocks = []
    for p in paragraphs or []:
        if p and p.strip():
            blocks.append({"type": "p", "text": p.strip()})
    if bullets:
        blocks.append({"type": "bullets", "items": [b for b in bullets if b and b.strip()]})
    if numbered:
        blocks.append({"type": "numbered", "items": [b for b in numbered if b and b.strip()]})
    if table:
        blocks.append({"type": "table", "rows": table})
    sec = {"heading": heading, "blocks": blocks}
    doc["sections"].append(sec)
    return sec


def validate_doc(doc) -> dict:
    """Coerce an arbitrary (possibly AI-produced) structure into the doc model."""
    if not isinstance(doc, dict):
        return new_doc("Document")
    out = new_doc(str(doc.get("title") or "Document")[:200], str(doc.get("subtitle") or "")[:300])
    for sec in doc.get("sections") or []:
        if not isinstance(sec, dict):
            continue
        blocks = []
        for b in sec.get("blocks") or []:
            if not isinstance(b, dict):
                continue
            t = b.get("type")
            if t in ("p", "quote", "h3") and b.get("text"):
                blocks.append({"type": t, "text": str(b["text"])[:6000]})
            elif t in ("bullets", "numbered") and isinstance(b.get("items"), list):
                items = [str(i)[:1500] for i in b["items"] if str(i).strip()][:40]
                if items:
                    blocks.append({"type": t, "items": items})
            elif t == "table" and isinstance(b.get("rows"), list):
                rows = [[str(c)[:300] for c in r][:12] for r in b["rows"] if isinstance(r, list)][:80]
                if rows:
                    blocks.append({"type": "table", "rows": rows})
        for key, typ in (("paragraphs", "p"), ("bullets", "bullets")):
            vals = sec.get(key)
            if isinstance(vals, list) and vals:
                if typ == "p":
                    blocks += [{"type": "p", "text": str(v)} for v in vals if str(v).strip()]
                else:
                    blocks.append({"type": "bullets", "items": [str(v) for v in vals if str(v).strip()]})
        if blocks:
            out["sections"].append({"heading": str(sec.get("heading") or "")[:200], "blocks": blocks})
    for src in doc.get("sources") or []:
        if isinstance(src, dict) and (src.get("title") or src.get("url")):
            out["sources"].append({k: str(src.get(k, ""))[:500] for k in ("title", "url", "license")})
    return out


def doc_to_markdown(doc: dict) -> str:
    lines = [f"# {doc.get('title', 'Document')}"]
    if doc.get("subtitle"):
        lines += ["", f"*{doc['subtitle']}*"]
    for sec in doc.get("sections", []):
        lines += ["", f"## {sec.get('heading', '')}"]
        for b in sec.get("blocks", []):
            lines.append("")
            t = b["type"]
            if t == "p":
                lines.append(b["text"])
            elif t == "h3":
                lines.append(f"### {b['text']}")
            elif t == "quote":
                lines.append("> " + b["text"])
            elif t == "image":
                lines.append(f"![{b.get('caption', '')}]({Path(b.get('path', '')).name})")
            elif t == "bullets":
                lines += [f"- {i}" for i in b["items"]]
            elif t == "numbered":
                lines += [f"{n}. {i}" for n, i in enumerate(b["items"], 1)]
            elif t == "table" and b["rows"]:
                rows = b["rows"]
                width = max(len(r) for r in rows)
                rows = [r + [""] * (width - len(r)) for r in rows]
                lines.append("| " + " | ".join(c.replace("|", "/") for c in rows[0]) + " |")
                lines.append("|" + "---|" * width)
                for r in rows[1:]:
                    lines.append("| " + " | ".join(c.replace("|", "/") for c in r) + " |")
    if doc.get("sources"):
        lines += ["", "## Sources"]
        for s in doc["sources"]:
            lic = f" — {s['license']}" if s.get("license") else ""
            lines.append(f"- {s.get('title', '')} {s.get('url', '')}{lic}".rstrip())
    return "\n".join(lines).strip() + "\n"


def markdown_to_doc(md: str, fallback_title: str = "Document") -> dict:
    doc = new_doc(fallback_title)
    sec = None
    para, items, kind, table = [], [], None, []

    def flush():
        nonlocal para, items, kind, table, sec
        if sec is None and (para or items or table):
            sec = {"heading": "", "blocks": []}
            doc["sections"].append(sec)
        if para:
            sec["blocks"].append({"type": "p", "text": " ".join(para).strip()})
        if items:
            sec["blocks"].append({"type": kind, "items": items})
        if table:
            rows = [r for r in table if not re.match(r"^\s*\|?\s*:?-{2,}", "|".join(r))]
            if rows:
                sec["blocks"].append({"type": "table", "rows": rows})
        para, items, kind, table = [], [], None, []

    for raw in (md or "").splitlines():
        line = raw.rstrip()
        s = line.strip()
        if re.match(r"^#\s+", s) and not doc["sections"] and doc["title"] == fallback_title:
            flush()
            doc["title"] = s.lstrip("#").strip()
        elif re.match(r"^#{2}\s+", s) or (re.match(r"^#\s+", s)):
            flush()
            sec = {"heading": s.lstrip("#").strip(), "blocks": []}
            doc["sections"].append(sec)
        elif re.match(r"^#{3,6}\s+", s):
            flush()
            if sec is None:
                sec = {"heading": "", "blocks": []}
                doc["sections"].append(sec)
            sec["blocks"].append({"type": "h3", "text": s.lstrip("#").strip()})
        elif re.match(r"^[-*•]\s+", s):
            if kind not in (None, "bullets") or para:
                flush()
            kind = "bullets"
            items.append(re.sub(r"^[-*•]\s+", "", s))
        elif re.match(r"^\d+[.)]\s+", s):
            if kind not in (None, "numbered") or para:
                flush()
            kind = "numbered"
            items.append(re.sub(r"^\d+[.)]\s+", "", s))
        elif s.startswith("|") and s.endswith("|"):
            if para or items:
                flush()
            table.append([c.strip() for c in s.strip("|").split("|")])
        elif s.startswith(">"):
            flush()
            if sec is None:
                sec = {"heading": "", "blocks": []}
                doc["sections"].append(sec)
            sec["blocks"].append({"type": "quote", "text": s.lstrip("> ").strip()})
        elif not s:
            flush()
        else:
            if items or table:
                flush()
            para.append(s)
    flush()
    doc["sections"] = [x for x in doc["sections"] if x["heading"] or x["blocks"]]
    if doc["sections"] and doc["sections"][0]["heading"] == "" and len(doc["sections"]) > 1:
        doc["sections"][0]["heading"] = "Introduction"
    return doc


def doc_plain_text(doc: dict) -> str:
    out = [doc.get("title", "")]
    for sec in doc.get("sections", []):
        out.append(sec.get("heading", ""))
        for b in sec.get("blocks", []):
            if b["type"] in ("p", "quote", "h3"):
                out.append(strip_md(b["text"]))
            elif b["type"] in ("bullets", "numbered"):
                out += [strip_md(i) + ("" if i.rstrip().endswith((".", "!", "?")) else ".") for i in b["items"]]
            elif b["type"] == "table":
                out += [", ".join(r) for r in b["rows"]]
    return "\n\n".join(x for x in out if x)


# ---------------------------------------------------------------- HTML -> text
class _TextExtractor(HTMLParser):
    SKIP = {"script", "style", "noscript", "svg", "nav", "footer", "header", "form", "aside", "iframe"}
    BLOCK = {"p", "div", "br", "li", "h1", "h2", "h3", "h4", "h5", "h6", "tr", "section", "article", "td"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts, self.title, self._skip, self._in_title = [], "", 0, False

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skip += 1
        if tag == "title":
            self._in_title = True
        if tag in self.BLOCK:
            self.parts.append("\n")
        if tag in ("h1", "h2", "h3") and not self._skip:
            self.parts.append("\n## ")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self._skip:
            self._skip -= 1
        if tag == "title":
            self._in_title = False
        if tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if self._in_title:
            self.title += data
        elif not self._skip:
            self.parts.append(data)


def html_to_text(raw_html: str):
    p = _TextExtractor()
    try:
        p.feed(raw_html)
    except Exception:  # malformed HTML: fall back to tag stripping
        return "", normalize_ws(html.unescape(re.sub(r"<[^>]+>", " ", raw_html)))
    text = "".join(p.parts)
    lines = [re.sub(r"\s+", " ", ln).strip() for ln in text.splitlines()]
    lines = [ln for ln in lines if ln and (len(ln.split()) > 3 or ln.startswith("## "))]
    cleaned = "\n".join(lines)
    cleaned = re.sub(r"^## (.+)$", r"\n## \1\n", cleaned, flags=re.M)
    return p.title.strip(), normalize_ws(cleaned)


# ---------------------------------------------------------------- file -> text
TEXT_EXT = {".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".log", ".py", ".js", ".ts", ".java",
            ".c", ".cpp", ".go", ".rs", ".rb", ".php", ".sql", ".yaml", ".yml", ".xml", ".ini", ".css"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif"}


def ocr_image(img, lang: str = "eng") -> str:
    try:
        import pytesseract
        return pytesseract.image_to_string(img, lang=lang)
    except Exception:
        return ""


def _mark_pdf_headings(page_text: str) -> str:
    """PDF text has no structure: turn short title-like lines into markdown headings and
    keep wrapped lines together, so sentences and sections survive extraction."""
    lines = [ln.strip() for ln in (page_text or "").splitlines()]
    out = []
    for i, ln in enumerate(lines):
        if not ln:
            out.append("")
            continue
        nxt = next((x for x in lines[i + 1:] if x), "")
        words_n = len(ln.split())
        title_like = (1 <= words_n <= 8 and not re.search(r"[.,;:!?]$", ln) and ln[:1].isupper()
                      and (not nxt or nxt[:1].isupper()) and not re.match(r"^[\d\W]+$", ln)
                      and sum(c.isalpha() for c in ln) >= 3)
        prev_open = out and out[-1] and not re.search(r"[.!?:]$", out[-1]) and not out[-1].startswith("## ")
        if title_like and not (prev_open and len(out[-1].split()) > 8):
            out += ["", "## " + ln, ""]
        else:
            out.append(ln)
    return re.sub(r"\n{3,}", "\n\n", "\n".join(out)).strip()


def file_to_text(path, max_chars: int = 400000, ocr: bool = True, lang: str = "eng") -> str:
    path = Path(path)
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            import pdfplumber
            out = []
            with pdfplumber.open(str(path)) as pdf:
                for i, page in enumerate(pdf.pages[:500]):
                    out.append(page.extract_text() or "")
                    if sum(map(len, out)) > max_chars:
                        break
            text = "\n\n".join(_mark_pdf_headings(t) for t in out)
            if ocr and len(text.strip()) < 40:  # likely scanned
                import pypdfium2 as pdfium
                pdf = pdfium.PdfDocument(str(path))
                pieces = []
                for i in range(min(len(pdf), 30)):
                    page = pdf[i]
                    w, h = page.get_size()
                    scale = min(2.0, (30_000_000 / max(1.0, w * h)) ** 0.5)
                    img = page.render(scale=scale).to_pil()
                    pieces.append(ocr_image(img, lang))
                    img.close()
                pdf.close()
                text = "\n\n".join(pieces)
            return text[:max_chars]
        if ext == ".docx":
            import docx
            d = docx.Document(str(path))
            parts = []
            for p in d.paragraphs:
                if p.style is not None and p.style.name.lower().startswith("heading") and p.text.strip():
                    parts.append("\n## " + p.text.strip() + "\n")
                else:
                    parts.append(p.text)
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            return normalize_ws("\n".join(parts))[:max_chars]
        if ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(str(path))
            parts = []
            for n, slide in enumerate(prs.slides, 1):
                parts.append(f"\n## Slide {n}\n")
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        parts.append(shape.text_frame.text)
            return normalize_ws("\n".join(parts))[:max_chars]
        if ext in (".xlsx", ".xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets:
                parts.append(f"\n## {ws.title}\n")
                for row in ws.iter_rows(values_only=True, max_row=2000, max_col=100):
                    if any(v is not None for v in row):
                        parts.append(", ".join("" if v is None else str(v) for v in row))
            wb.close()
            return "\n".join(parts)[:max_chars]
        if ext in (".html", ".htm"):
            _, text = html_to_text(path.read_text(encoding="utf-8", errors="replace"))
            return text[:max_chars]
        if ext in IMAGE_EXT and ocr:
            from PIL import Image
            with Image.open(path) as img:
                return ocr_image(img.convert("RGB"), lang)[:max_chars]
        if ext in TEXT_EXT or ext == "":
            return path.read_text(encoding="utf-8", errors="replace")[:max_chars]
    except Exception as exc:  # corrupted or unsupported file
        return f"[Could not read {path.name}: {type(exc).__name__}]"
    return ""


def read_table_file(path):
    """Return (header, rows) from CSV/TSV/XLSX."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True, max_row=100000, max_col=200)]
        wb.close()
    else:
        raw = path.read_text(encoding="utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(raw[:5000], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = []
        for i, r in enumerate(csv.reader(io.StringIO(raw), dialect)):
            if i >= 100000:
                break
            rows.append(r[:200])
    rows = [r for r in rows if r and any(c not in (None, "") for c in r)]
    if not rows:
        return [], []
    header = [str(c) if c not in (None, "") else f"Column {i + 1}" for i, c in enumerate(rows[0])]
    return header, rows[1:]
